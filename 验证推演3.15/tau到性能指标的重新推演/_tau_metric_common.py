#!/usr/bin/env python3
"""Shared helpers for tau-to-metrics validation scripts."""

from __future__ import annotations

import csv
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


LOCK_RE = re.compile(r"lock\s*=\s*(-?\d+)", re.IGNORECASE)
CH_RE = re.compile(r"ch\s*=\s*(-?\d+)", re.IGNORECASE)


@dataclass(frozen=True)
class MatrixData:
    row_labels: list[int]
    col_labels: list[int]
    values: list[list[float]]

    @property
    def shape(self) -> tuple[int, int]:
        return len(self.row_labels), len(self.col_labels)


@dataclass(frozen=True)
class ComparisonStats:
    matrix_mae_pp: float
    diag_mae_pp: float
    offdiag_mae_pp: float
    top1_match: int
    top2_match: int
    columns: int


def script_dir() -> Path:
    return Path(__file__).resolve().parent


def repo_root() -> Path:
    return script_dir().parents[1]


def default_final_data_dir() -> Path:
    return repo_root() / "数据内容" / "对比度数据" / "最终数据"


def default_theory_csv() -> Path:
    return script_dir() / "theory_probability_matrix.csv"


def default_output_dir(path: str | None) -> Path:
    out_dir = Path(path) if path else script_dir()
    out_dir.mkdir(parents=True, exist_ok=True)
    return out_dir


def mode_labels(lmin: int, lmax: int) -> list[int]:
    return list(range(lmin, lmax + 1))


def parse_float(value: object) -> float:
    if value is None:
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return float("nan")
    return float(text)


def k_from_lr(length_mm: float, radius_mm: float) -> float:
    ratio = float(length_mm) / float(radius_mm)
    if not (0.0 < ratio < 1.0):
        raise ValueError(f"L/R must lie in (0, 1), got {ratio!r}")
    return math.acos(math.sqrt(1.0 - ratio)) / math.pi


def circular_distance(a: float, b: float) -> float:
    diff = abs(float(a) - float(b))
    return diff if diff <= 0.5 else 1.0 - diff


def positions_for_modes(l_values: Iterable[int], k_value: float) -> dict[int, float]:
    return {l_value: (float(l_value) * k_value) % 1.0 for l_value in l_values}


def build_lorentz_probability_matrix(
    l_values: list[int], finesse: float, k_value: float
) -> tuple[MatrixData, dict[tuple[int, int], tuple[float, float, float]], float, float]:
    positions = positions_for_modes(l_values, k_value)
    row_labels = list(l_values)
    col_labels = list(l_values)
    matrix = [[0.0 for _ in col_labels] for _ in row_labels]
    pairwise: dict[tuple[int, int], tuple[float, float, float]] = {}
    s_min = float("inf")

    for j, lock_l in enumerate(col_labels):
        col_sum = 0.0
        for i, det_l in enumerate(row_labels):
            if det_l == lock_l:
                s_ij = 0.0
                tau_ij = float("inf")
                kernel = 1.0
            else:
                s_ij = circular_distance(positions[det_l], positions[lock_l])
                tau_ij = finesse * s_ij
                kernel = 1.0 / (1.0 + 4.0 * tau_ij * tau_ij)
                s_min = min(s_min, s_ij)
            matrix[i][j] = kernel
            pairwise[(det_l, lock_l)] = (s_ij, tau_ij, kernel)
            col_sum += kernel

        for i in range(len(row_labels)):
            matrix[i][j] /= col_sum

    tau_eff = finesse * s_min
    return MatrixData(row_labels=row_labels, col_labels=col_labels, values=matrix), pairwise, s_min, tau_eff


def column_sums(matrix: MatrixData) -> list[float]:
    return [
        sum(matrix.values[row_idx][col_idx] for row_idx in range(len(matrix.row_labels)))
        for col_idx in range(len(matrix.col_labels))
    ]


def read_matrix_sheet(path: Path, sheet_name: str) -> MatrixData:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise KeyError(f"{path} is missing sheet {sheet_name!r}")
    worksheet = workbook[sheet_name]

    col_labels: list[int] = []
    for col_idx in range(2, worksheet.max_column + 1):
        header = worksheet.cell(1, col_idx).value
        match = LOCK_RE.search(str(header)) if header is not None else None
        if match:
            col_labels.append(int(match.group(1)))

    row_labels: list[int] = []
    rows: list[list[float]] = []
    started = False
    for row_idx in range(2, worksheet.max_row + 1):
        label = worksheet.cell(row_idx, 1).value
        match = CH_RE.search(str(label)) if label is not None else None
        if not match:
            if started:
                break
            continue
        started = True
        row_labels.append(int(match.group(1)))
        rows.append(
            [parse_float(worksheet.cell(row_idx, col_idx).value) for col_idx in range(2, 2 + len(col_labels))]
        )

    workbook.close()

    if not row_labels or not col_labels:
        raise ValueError(f"{path}:{sheet_name} does not contain a valid matrix block")

    return MatrixData(row_labels=row_labels, col_labels=col_labels, values=rows)


def normalize_columns(matrix: MatrixData, scale: float = 1.0) -> MatrixData:
    out = [[0.0 for _ in matrix.col_labels] for _ in matrix.row_labels]
    for col_idx in range(len(matrix.col_labels)):
        col_sum = sum(max(0.0, float(matrix.values[row_idx][col_idx])) for row_idx in range(len(matrix.row_labels)))
        if col_sum <= 0.0:
            raise ValueError(f"Column {matrix.col_labels[col_idx]} has non-positive sum")
        for row_idx in range(len(matrix.row_labels)):
            out[row_idx][col_idx] = max(0.0, float(matrix.values[row_idx][col_idx])) * scale / col_sum
    return MatrixData(row_labels=matrix.row_labels, col_labels=matrix.col_labels, values=out)


def read_theory_matrix_csv(path: Path) -> MatrixData:
    with path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        col_labels = [int(item.split("=")[-1]) for item in header[1:]]
        row_labels: list[int] = []
        rows: list[list[float]] = []
        for row in reader:
            row_labels.append(int(row[0].split("=")[-1]))
            rows.append([float(cell) for cell in row[1:]])
    return MatrixData(row_labels=row_labels, col_labels=col_labels, values=rows)


def write_matrix_csv(path: Path, matrix: MatrixData) -> None:
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["detector_l"] + [f"lock_l={label}" for label in matrix.col_labels])
        for row_idx, label in enumerate(matrix.row_labels):
            writer.writerow([f"detector_l={label}"] + [f"{value:.12f}" for value in matrix.values[row_idx]])


def compare_matrices(theory: MatrixData, experiment: MatrixData) -> ComparisonStats:
    assert_same_axes(theory, experiment)
    total_error = 0.0
    diag_error = 0.0
    offdiag_error = 0.0
    total_cells = len(theory.row_labels) * len(theory.col_labels)
    offdiag_cells = total_cells - len(theory.row_labels)
    top1_match = 0
    top2_match = 0

    for col_idx in range(len(theory.col_labels)):
        theory_rank: list[tuple[int, float]] = []
        exp_rank: list[tuple[int, float]] = []
        for row_idx in range(len(theory.row_labels)):
            diff = abs(theory.values[row_idx][col_idx] - experiment.values[row_idx][col_idx])
            total_error += diff
            if row_idx == col_idx:
                diag_error += diff
            else:
                offdiag_error += diff
                theory_rank.append((theory.row_labels[row_idx], theory.values[row_idx][col_idx]))
                exp_rank.append((experiment.row_labels[row_idx], experiment.values[row_idx][col_idx]))

        theory_rank.sort(key=lambda item: item[1], reverse=True)
        exp_rank.sort(key=lambda item: item[1], reverse=True)
        if theory_rank[0][0] == exp_rank[0][0]:
            top1_match += 1
        if sorted((theory_rank[0][0], theory_rank[1][0])) == sorted((exp_rank[0][0], exp_rank[1][0])):
            top2_match += 1

    return ComparisonStats(
        matrix_mae_pp=100.0 * total_error / total_cells,
        diag_mae_pp=100.0 * diag_error / len(theory.row_labels),
        offdiag_mae_pp=100.0 * offdiag_error / offdiag_cells,
        top1_match=top1_match,
        top2_match=top2_match,
        columns=len(theory.col_labels),
    )


def assert_same_axes(left: MatrixData, right: MatrixData) -> None:
    if left.row_labels != right.row_labels or left.col_labels != right.col_labels:
        raise ValueError(
            "Matrix axis mismatch: "
            f"left rows={left.row_labels}, cols={left.col_labels}; "
            f"right rows={right.row_labels}, cols={right.col_labels}"
        )


def per_channel_metrics(matrix: MatrixData) -> list[dict[str, float]]:
    metrics: list[dict[str, float]] = []
    for idx, label in enumerate(matrix.col_labels):
        diag = matrix.values[idx][idx]
        offdiag = [matrix.values[row_idx][idx] for row_idx in range(len(matrix.row_labels)) if row_idx != idx]
        leak_sum = sum(offdiag)
        leak_max = max(offdiag)
        er_sum = 10.0 * math.log10(diag / leak_sum) if leak_sum > 0.0 else float("inf")
        er_max = 10.0 * math.log10(diag / leak_max) if leak_max > 0.0 else float("inf")
        metrics.append(
            {
                "l": float(label),
                "diag_share": diag,
                "e_sort": 1.0 - diag,
                "er_sum_db": er_sum,
                "er_max_db": er_max,
            }
        )
    return metrics


def eta_sort(matrix: MatrixData) -> float:
    return sum(matrix.values[idx][idx] for idx in range(len(matrix.row_labels))) / len(matrix.row_labels)


def mutual_information_bits(matrix: MatrixData) -> float:
    d = len(matrix.col_labels)
    prior = 1.0 / d
    p_out = [sum(matrix.values[row_idx][col_idx] for col_idx in range(d)) * prior for row_idx in range(d)]
    info = 0.0
    for col_idx in range(d):
        for row_idx in range(d):
            p = matrix.values[row_idx][col_idx]
            if p > 0.0 and p_out[row_idx] > 0.0:
                info += prior * p * math.log2(p / p_out[row_idx])
    return info


def read_er_summary_sheet(path: Path) -> dict[int, dict[str, float]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook["ER_summary"]
    header = [worksheet.cell(1, col_idx).value for col_idx in range(1, worksheet.max_column + 1)]
    header_map = {str(value): idx for idx, value in enumerate(header)}
    required = ["ER_sum_mean_dB", "ER_max_mean_dB"]
    for key in required:
        if key not in header_map:
            workbook.close()
            raise KeyError(f"{path}: ER_summary is missing {key!r}")

    results: dict[int, dict[str, float]] = {}
    row_idx = 2
    while True:
        channel_value = worksheet.cell(row_idx, 1).value
        if channel_value is None:
            break
        channel = int(channel_value)
        results[channel] = {
            "er_sum_mean_db": parse_float(worksheet.cell(row_idx, header_map["ER_sum_mean_dB"] + 1).value),
            "er_max_mean_db": parse_float(worksheet.cell(row_idx, header_map["ER_max_mean_dB"] + 1).value),
        }
        row_idx += 1

    workbook.close()
    return results


def format_float(value: float, digits: int = 6) -> str:
    if math.isinf(value):
        return "inf" if value > 0.0 else "-inf"
    if math.isnan(value):
        return "nan"
    return f"{value:.{digits}f}"
