#!/usr/bin/env python3
"""Read-only Q1 verification for the per-lock-column model.

Goal:
1. Verify that the trusted workbooks are naturally organized by lock columns.
2. Verify the per-lock model y^(j) ~= S x^(j) using existing de-embedded X.
3. Verify that percentage heatmaps and ER metrics are consistent with X-column
   normalization rather than a mixed measurement mouth.

This script never edits the original data files. All outputs are written next
to the script itself.
"""

from __future__ import annotations

import json
import math
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


LOCK_RE = re.compile(r"lock\s*=\s*(\d+)", re.IGNORECASE)
CH_RE = re.compile(r"ch\s*=\s*(\d+)", re.IGNORECASE)


@dataclass
class Axes:
    channels: list[int]
    locks: list[int]


@dataclass
class RunData:
    run_name: str
    axes: Axes
    s_matrix: np.ndarray
    x_matrix: np.ndarray
    y_matrix: np.ndarray


def parse_float(value) -> float:
    if value is None:
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("_x000d_", "")
    if text in {"", "NaN", "nan"}:
        return float("nan")
    if text == "Inf":
        return float("inf")
    if text == "-Inf":
        return float("-inf")
    return float(text)


def parse_matrix_sheet(path: Path, sheet_name: str) -> tuple[Axes, np.ndarray]:
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise KeyError(f"{path.name} missing sheet {sheet_name}")
    ws = wb[sheet_name]

    locks: list[int] = []
    for col in range(2, ws.max_column + 1):
        header = ws.cell(1, col).value
        if header is None:
            continue
        match = LOCK_RE.search(str(header))
        if match:
            locks.append(int(match.group(1)))

    channels: list[int] = []
    rows: list[list[float]] = []
    for row in range(2, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if label is None:
            continue
        match = CH_RE.search(str(label))
        if not match:
            continue
        channels.append(int(match.group(1)))
        rows.append([parse_float(ws.cell(row, col).value) for col in range(2, 2 + len(locks))])

    wb.close()
    if not channels or not locks:
        raise ValueError(f"{path.name}:{sheet_name} has no matrix axes")
    return Axes(channels=channels, locks=locks), np.asarray(rows, dtype=float)


def safe_db10(num: float, den: float) -> float:
    if not math.isfinite(num) or not math.isfinite(den) or num <= 0 or den <= 0:
        return float("nan")
    return 10.0 * math.log10(num / den)


def compute_percentage(x_matrix: np.ndarray) -> np.ndarray:
    col_sums = np.sum(x_matrix, axis=0, keepdims=True)
    return np.divide(x_matrix, col_sums, out=np.zeros_like(x_matrix), where=col_sums > 0)


def compute_er_metrics(x_matrix: np.ndarray, axes: Axes) -> tuple[dict[int, float], dict[int, float], dict[int, float]]:
    ch_to_row = {ch: idx for idx, ch in enumerate(axes.channels)}
    er_sum: dict[int, float] = {}
    er_max: dict[int, float] = {}
    signal_share: dict[int, float] = {}

    for col, lock in enumerate(axes.locks):
        if lock not in ch_to_row:
            continue
        sig_row = ch_to_row[lock]
        col_vec = np.clip(x_matrix[:, col], 0.0, None)
        signal = float(col_vec[sig_row])
        leak = np.delete(col_vec, sig_row)
        leak_sum = float(np.sum(leak)) if leak.size else float("nan")
        leak_max = float(np.max(leak)) if leak.size else float("nan")
        total = float(np.sum(col_vec))
        er_sum[lock] = safe_db10(signal, leak_sum)
        er_max[lock] = safe_db10(signal, leak_max)
        signal_share[lock] = signal / total if total > 0 else float("nan")

    return er_sum, er_max, signal_share


def load_run(final_dir: Path, run_name: str) -> RunData:
    run_dir = final_dir / run_name
    axes_s, s_matrix = parse_matrix_sheet(run_dir / "calib_matrix_with_uncertainty.xlsx", "mean")
    axes_x, x_matrix = parse_matrix_sheet(run_dir / "full9_matrix_with_uncertainty.xlsx", "mean")
    axes_y, y_matrix = parse_matrix_sheet(run_dir / "full9_matrix_with_uncertainty.xlsx", "raw_y_mean")

    if axes_s != axes_x or axes_s != axes_y:
        raise ValueError(f"Axis mismatch in run {run_name}")

    return RunData(
        run_name=run_name,
        axes=axes_s,
        s_matrix=s_matrix,
        x_matrix=x_matrix,
        y_matrix=y_matrix,
    )


def parse_aggregate_matrix(path: Path, sheet_name: str) -> tuple[Axes, np.ndarray]:
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise KeyError(f"{path.name} missing sheet {sheet_name}")
    ws = wb[sheet_name]

    locks: list[int] = []
    for col in range(2, 11):
        header = ws.cell(1, col).value
        if header is None:
            continue
        match = LOCK_RE.search(str(header))
        if match:
            locks.append(int(match.group(1)))

    channels: list[int] = []
    rows: list[list[float]] = []
    for row in range(2, 11):
        label = ws.cell(row, 1).value
        if label is None:
            continue
        match = CH_RE.search(str(label))
        if not match:
            continue
        channels.append(int(match.group(1)))
        rows.append([parse_float(ws.cell(row, col).value) for col in range(2, 11)])

    wb.close()
    return Axes(channels=channels, locks=locks), np.asarray(rows, dtype=float)


def parse_er_summary(path: Path) -> dict[int, dict[str, float]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["ER_summary"]
    headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
    out: dict[int, dict[str, float]] = {}
    for row in range(2, ws.max_row + 1):
        channel = parse_float(ws.cell(row, 1).value)
        if not math.isfinite(channel):
            continue
        item: dict[str, float] = {}
        for col in range(2, ws.max_column + 1):
            key = str(headers[col - 1])
            item[key] = parse_float(ws.cell(row, col).value)
        out[int(channel)] = item
    wb.close()
    return out


def max_abs_diff(a: np.ndarray, b: np.ndarray) -> float:
    return float(np.max(np.abs(a - b)))


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    data_root = script_dir.parent
    final_dir = data_root / "最终数据"
    aggregate_path = final_dir / "triplicate_full9_aggregate.xlsx"

    runs = [load_run(final_dir, run_name) for run_name in ["1", "2", "3"]]
    axes = runs[0].axes

    results: dict[str, object] = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "script": str(Path(__file__).resolve()),
        "data_root": str(data_root),
        "final_dir": str(final_dir),
        "axes": {
            "channels": axes.channels,
            "locks": axes.locks,
            "shape_interpretation": {
                "single_lock_column": "9x1",
                "all_locks_matrix": "9x9",
            },
        },
        "runs": {},
        "aggregate_checks": {},
    }

    computed_pct_by_run: dict[str, np.ndarray] = {}
    er_sum_by_run: dict[str, dict[int, float]] = {}
    er_max_by_run: dict[str, dict[int, float]] = {}
    share_by_run: dict[str, dict[int, float]] = {}

    for run in runs:
        y_hat = run.s_matrix @ run.x_matrix
        diff = y_hat - run.y_matrix
        global_rel = float(np.linalg.norm(diff) / np.linalg.norm(run.y_matrix))

        er_sum, er_max, signal_share = compute_er_metrics(run.x_matrix, run.axes)
        pct = compute_percentage(run.x_matrix) * 100.0

        per_lock = []
        for col, lock in enumerate(run.axes.locks):
            y_col = run.y_matrix[:, col]
            d_col = diff[:, col]
            rel_l2 = float(np.linalg.norm(d_col) / np.linalg.norm(y_col))
            abs_l2 = float(np.linalg.norm(d_col))
            max_abs = float(np.max(np.abs(d_col)))
            diag_share = float(signal_share[lock] * 100.0)
            per_lock.append(
                {
                    "lock": lock,
                    "rel_l2_residual": rel_l2,
                    "abs_l2_residual_W": abs_l2,
                    "max_abs_residual_W": max_abs,
                    "diag_share_percent": diag_share,
                    "x_diag_uW": float(run.x_matrix[col, col] * 1e6),
                    "y_diag_uW": float(run.y_matrix[col, col] * 1e6),
                    "er_sum_db": float(er_sum[lock]),
                    "er_max_db": float(er_max[lock]),
                }
            )

        results["runs"][run.run_name] = {
            "row_labels": [f"ch={ch}" for ch in run.axes.channels],
            "column_labels": [f"lock={lk}" for lk in run.axes.locks],
            "global_rel_l2_residual": global_rel,
            "max_lock_rel_l2_residual": max(item["rel_l2_residual"] for item in per_lock),
            "worst_lock": max(per_lock, key=lambda item: item["rel_l2_residual"])["lock"],
            "per_lock": per_lock,
        }

        computed_pct_by_run[run.run_name] = pct
        er_sum_by_run[run.run_name] = er_sum
        er_max_by_run[run.run_name] = er_max
        share_by_run[run.run_name] = signal_share

    axes_pct, pct_sheet_1 = parse_aggregate_matrix(aggregate_path, "1_percentage")
    _, pct_sheet_2 = parse_aggregate_matrix(aggregate_path, "2_percentage")
    _, pct_sheet_3 = parse_aggregate_matrix(aggregate_path, "3_percentage")
    _, pct_sheet_combined = parse_aggregate_matrix(aggregate_path, "combined_percentage")
    _, x_sheet_combined = parse_aggregate_matrix(aggregate_path, "combined_deembedded_uW")
    er_summary = parse_er_summary(aggregate_path)

    if axes_pct != axes:
        raise ValueError("Aggregate workbook axes mismatch")

    pct_run_diffs = {
        "1": max_abs_diff(computed_pct_by_run["1"], pct_sheet_1),
        "2": max_abs_diff(computed_pct_by_run["2"], pct_sheet_2),
        "3": max_abs_diff(computed_pct_by_run["3"], pct_sheet_3),
    }
    pct_combined_mean = (computed_pct_by_run["1"] + computed_pct_by_run["2"] + computed_pct_by_run["3"]) / 3.0
    pct_combined_diff = max_abs_diff(pct_combined_mean, pct_sheet_combined)

    x_combined_from_runs = sum(run.x_matrix for run in runs) / 3.0 * 1e6
    x_combined_diff = max_abs_diff(x_combined_from_runs, x_sheet_combined)

    er_checks: list[dict[str, float]] = []
    max_er_sum_diff = 0.0
    max_er_max_diff = 0.0
    max_share_diff = 0.0
    for lock in axes.locks:
        mean_er_sum = float(np.mean([er_sum_by_run[r][lock] for r in ["1", "2", "3"]]))
        mean_er_max = float(np.mean([er_max_by_run[r][lock] for r in ["1", "2", "3"]]))
        mean_share = float(np.mean([share_by_run[r][lock] for r in ["1", "2", "3"]]) * 100.0)
        row = er_summary[lock]
        diff_sum = abs(mean_er_sum - row["ER_sum_mean_dB"])
        diff_max = abs(mean_er_max - row["ER_max_mean_dB"])
        diff_share = abs(mean_share - row["SignalShare_mean_%"])
        max_er_sum_diff = max(max_er_sum_diff, diff_sum)
        max_er_max_diff = max(max_er_max_diff, diff_max)
        max_share_diff = max(max_share_diff, diff_share)
        er_checks.append(
            {
                "lock": lock,
                "recomputed_ER_sum_mean_dB": mean_er_sum,
                "aggregate_ER_sum_mean_dB": row["ER_sum_mean_dB"],
                "abs_diff_ER_sum_dB": diff_sum,
                "recomputed_ER_max_mean_dB": mean_er_max,
                "aggregate_ER_max_mean_dB": row["ER_max_mean_dB"],
                "abs_diff_ER_max_dB": diff_max,
                "recomputed_signal_share_mean_percent": mean_share,
                "aggregate_signal_share_mean_percent": row["SignalShare_mean_%"],
                "abs_diff_signal_share_percent": diff_share,
            }
        )

    results["aggregate_checks"] = {
        "percentage_sheet_max_abs_diff_percent": pct_run_diffs,
        "combined_percentage_max_abs_diff_percent": pct_combined_diff,
        "combined_deembedded_uW_max_abs_diff_uW": x_combined_diff,
        "er_summary_consistency": {
            "max_abs_diff_ER_sum_dB": max_er_sum_diff,
            "max_abs_diff_ER_max_dB": max_er_max_diff,
            "max_abs_diff_signal_share_percent": max_share_diff,
            "per_lock": er_checks,
        },
    }

    json_path = script_dir / "q1_per_lock_model_results.json"
    md_path = script_dir / "q1_per_lock_model_results.md"
    with json_path.open("w", encoding="utf-8") as fh:
        json.dump(results, fh, ensure_ascii=False, indent=2)

    lines: list[str] = []
    lines.append("# Q1 逐锁频列向量模型验算结果")
    lines.append("")
    lines.append(f"- 生成时间：`{results['generated_at']}`")
    lines.append(f"- 数据目录：`{final_dir}`")
    lines.append(f"- 锁定态维度：单列 `9x1`，整组 `9x9`")
    lines.append("")
    lines.append("## 逐运行残差")
    lines.append("")
    for run_name in ["1", "2", "3"]:
        run_result = results["runs"][run_name]
        lines.append(f"### Run {run_name}")
        lines.append(f"- 全局相对残差：`{run_result['global_rel_l2_residual']:.6f}`")
        lines.append(f"- 最差列：`lock={run_result['worst_lock']}`")
        lines.append(f"- 最差列相对残差：`{run_result['max_lock_rel_l2_residual']:.6f}`")
        lines.append("")
        lines.append("| lock | rel_L2_residual | diag_share(%) | ER_sum(dB) | ER_max(dB) |")
        lines.append("| --- | ---: | ---: | ---: | ---: |")
        for item in run_result["per_lock"]:
            lines.append(
                f"| {item['lock']} | {item['rel_l2_residual']:.6f} | "
                f"{item['diag_share_percent']:.3f} | {item['er_sum_db']:.3f} | {item['er_max_db']:.3f} |"
            )
        lines.append("")

    lines.append("## 聚合口径一致性")
    lines.append("")
    lines.append("- 各运行百分比热图与由 `X` 列归一化直接重算结果的最大绝对差（百分点）：")
    for run_name in ["1", "2", "3"]:
        lines.append(
            f"  - Run {run_name}: `{pct_run_diffs[run_name]:.6f}`"
        )
    lines.append(f"- 合并百分比热图最大绝对差（百分点）：`{pct_combined_diff:.6f}`")
    lines.append(f"- 合并去嵌入矩阵最大绝对差（uW）：`{x_combined_diff:.6f}`")
    lines.append(f"- ER_sum 聚合表最大绝对差（dB）：`{max_er_sum_diff:.6f}`")
    lines.append(f"- ER_max 聚合表最大绝对差（dB）：`{max_er_max_diff:.6f}`")
    lines.append(f"- 对角占比聚合表最大绝对差（百分点）：`{max_share_diff:.6f}`")
    lines.append("")
    lines.append("## 用途边界")
    lines.append("")
    lines.append("- 本结果仅用于 Q1：证明数据结构和计算口径本质上是逐锁频列向量/列矩阵。")
    lines.append("- 本结果不自动证明标定条件向使用条件迁移的物理充分性，因此不能替代 Q2。")
    lines.append("- 若需对审稿回复使用，建议优先引用“列向量定义”和“口径一致性”，谨慎使用最差列残差。")
    lines.append("")

    with md_path.open("w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")

    print(f"Wrote {json_path}")
    print(f"Wrote {md_path}")


if __name__ == "__main__":
    main()
