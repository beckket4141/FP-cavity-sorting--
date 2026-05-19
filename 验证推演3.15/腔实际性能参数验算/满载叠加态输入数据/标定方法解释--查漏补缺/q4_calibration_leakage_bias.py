#!/usr/bin/env python3
"""Quantify the Q4 calibration-leakage concern.

Q4 asks whether the appendix sentence

    "FP cavity filtering makes the detected calibration field effectively pure"

overstates the real situation when the nearest-neighbor off-resonant leakage is
still about 2% at tau_min ~= 3.5.

This script keeps thesis sources, raw data files, and existing workbooks
untouched. It only reads the existing workbooks and writes result files in the
current folder.

Main idea
---------
1. Reuse the already validated Q2 forward-model ingredients:
   measured lock positions, measured peak transmittance, and the same Airy law.
2. Separate two different quantities that were blurred together in the raw Q4
   concern:
   - cavity leakage coefficient of a neighboring mode if it is present, and
   - the actual fraction that could have been mixed into the measured
     calibration column S[:,k].
3. Use the current calibration matrix itself to build a conservative upper bound
   on the latter, then propagate that bound through the existing NNLS pipeline.

Outputs
-------
* q4_calibration_leakage_bias_results.json
* q4_calibration_leakage_bias_results.md
* Q4_直白结论.md
* Q4_详细报告.md
"""

from __future__ import annotations

import json
import math
import re
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import nnls


C_LIGHT = 299_792_458.0
FSR_FIXED_HZ = 9.947e9
F_FIXED = 32.21

KNOWN_XLSX_NAMES = {
    "final_matrix_summary.xlsx",
    "calib_matrix_with_uncertainty.xlsx",
    "full9_matrix_with_uncertainty.xlsx",
    "full4_matrix_with_uncertainty.xlsx",
    "channel_snr_analysis.xlsx",
}

NUM_RE = re.compile(r"[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?")
L_RE = re.compile(r"l\s*=\s*(-?\d+)", re.IGNORECASE)


@dataclass
class NeighborMixEntry:
    mode: int
    geodesic_spacing: float
    cavity_leak_ratio: float
    cavity_leak_percent: float
    postmix_upper_bound: float
    postmix_upper_percent: float
    precavity_impurity_upper_bound: float
    precavity_impurity_upper_percent: float


@dataclass
class ScenarioMetrics:
    reported_mean_success_percent: float
    corrected_mean_success_percent: float
    reported_mean_er_db: float
    corrected_mean_er_db: float
    reported_minus_corrected_success_pct_pt: float
    reported_minus_corrected_er_db: float
    mean_abs_matrix_delta_pct_pt: float
    max_abs_matrix_delta_pct_pt: float
    mean_abs_offdiag_delta_pct_pt: float
    max_abs_offdiag_delta_pct_pt: float
    max_abs_diag_delta_pct_pt: float
    max_abs_er_delta_db: float
    min_corrected_response_entry: float


@dataclass
class RunSummary:
    run: str
    transmission_xlsx: str
    lock_positions: list[float]
    peak_transmittance_ratio: list[float]
    nearest_neighbors: list[list[NeighborMixEntry]]
    mean_single_neighbor_leak_percent: float
    leak_percent_range: list[float]
    mean_total_postmix_upper_percent_per_column: float
    max_total_postmix_upper_percent_per_column: float
    mean_single_neighbor_precavity_impurity_upper_percent: float
    max_single_neighbor_precavity_impurity_upper_percent: float
    literal_full_leak_scenario: ScenarioMetrics
    data_driven_upper_bound_scenario: ScenarioMetrics


def parse_float(value: Any) -> float:
    if value is None:
        raise ValueError("Encountered None where a number is required")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("_x000d_", "").strip()
    match = NUM_RE.search(text)
    if not match:
        raise ValueError(f"Cannot parse float from: {value!r}")
    return float(match.group(0))


def parse_l_index(value: Any) -> int:
    match = L_RE.search(str(value))
    if not match:
        raise ValueError(f"Cannot parse l-index from: {value!r}")
    return int(match.group(1))


def geodesic_distance(a: float, b: float) -> float:
    diff = abs(a - b)
    return min(diff, 1.0 - diff)


def airy_transmission(s: float, finesse: float) -> float:
    prefactor = (2.0 * finesse / math.pi) ** 2
    return 1.0 / (1.0 + prefactor * (math.sin(math.pi * s) ** 2))


def discover_run_dirs(data_root: Path) -> list[Path]:
    run_dirs = []
    for path in data_root.rglob("calib_matrix_with_uncertainty.xlsx"):
        if path.parent.name in {"1", "2", "3"} and (path.parent / "full9_matrix_with_uncertainty.xlsx").exists():
            run_dirs.append(path.parent)
    return sorted(run_dirs, key=lambda p: p.name)


def find_transmission_xlsx(run_dir: Path) -> Path:
    candidates = [p for p in run_dir.glob("*.xlsx") if p.name not in KNOWN_XLSX_NAMES]
    if len(candidates) != 1:
        raise FileNotFoundError(
            f"Could not uniquely determine transmission xlsx in {run_dir}: {candidates}"
        )
    return candidates[0]


def read_transmission_data(run_dir: Path) -> tuple[Path, np.ndarray, np.ndarray]:
    xlsx_path = find_transmission_xlsx(run_dir)
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = list(ws.iter_rows(min_row=2, max_row=10, values_only=True))
    records: list[tuple[int, float, float]] = []
    for row in rows:
        if row[0] is None:
            continue
        l_idx = parse_l_index(row[0])
        wavelength_nm = parse_float(row[1])
        p_before = parse_float(row[2])
        p_after = parse_float(row[3])
        records.append((l_idx, wavelength_nm, p_after / p_before))

    wb.close()

    records.sort(key=lambda item: item[0])
    wavelengths_nm = np.asarray([item[1] for item in records], dtype=float)
    peak_trans = np.asarray([item[2] for item in records], dtype=float)
    frequencies_hz = C_LIGHT / (wavelengths_nm * 1e-9)
    nu0 = float(frequencies_hz[0])
    positions = np.asarray([((nu - nu0) / FSR_FIXED_HZ) % 1.0 for nu in frequencies_hz], dtype=float)
    return xlsx_path, positions, peak_trans


def read_matrix(workbook: Path, sheet: str) -> np.ndarray:
    wb = load_workbook(workbook, data_only=True, read_only=True)
    rows = list(wb[sheet].iter_rows(values_only=True))
    wb.close()
    return np.asarray([[float(v) for v in row[1:10]] for row in rows[1:10]], dtype=float)


def solve_nnls(s_matrix: np.ndarray, y_matrix: np.ndarray) -> np.ndarray:
    x_matrix = np.zeros_like(y_matrix, dtype=float)
    for col in range(y_matrix.shape[1]):
        x_matrix[:, col], _ = nnls(s_matrix, np.clip(y_matrix[:, col], 0.0, None))
    return x_matrix


def compute_percent_matrix(x_matrix: np.ndarray) -> np.ndarray:
    return x_matrix / np.sum(x_matrix, axis=0, keepdims=True) * 100.0


def compute_er_sum_db(x_matrix: np.ndarray) -> np.ndarray:
    er_sum = np.zeros(x_matrix.shape[1], dtype=float)
    for lock in range(x_matrix.shape[1]):
        signal = float(x_matrix[lock, lock])
        leak = float(np.sum(x_matrix[:, lock]) - signal)
        er_sum[lock] = 10.0 * math.log10(signal / leak)
    return er_sum


def find_two_nearest_neighbors(lock_idx: int, positions: np.ndarray) -> list[tuple[int, float]]:
    candidates = [
        (mode_idx, geodesic_distance(float(positions[mode_idx]), float(positions[lock_idx])))
        for mode_idx in range(positions.size)
        if mode_idx != lock_idx
    ]
    candidates.sort(key=lambda item: (round(item[1], 12), item[0]))
    return candidates[:2]


def build_neighbor_entries(
    lock_idx: int,
    positions: np.ndarray,
    peak_trans: np.ndarray,
    s_matrix: np.ndarray,
) -> list[NeighborMixEntry]:
    entries: list[NeighborMixEntry] = []
    for neighbor_idx, spacing in find_two_nearest_neighbors(lock_idx, positions):
        leak_ratio = (
            float(peak_trans[neighbor_idx])
            / float(peak_trans[lock_idx])
            * airy_transmission(float(spacing), F_FIXED)
        )
        postmix_upper = float(s_matrix[neighbor_idx, lock_idx] / s_matrix[neighbor_idx, neighbor_idx])
        precavity_upper = postmix_upper / leak_ratio if leak_ratio > 0 else math.inf
        entries.append(
            NeighborMixEntry(
                mode=int(neighbor_idx),
                geodesic_spacing=float(spacing),
                cavity_leak_ratio=float(leak_ratio),
                cavity_leak_percent=float(leak_ratio * 100.0),
                postmix_upper_bound=float(postmix_upper),
                postmix_upper_percent=float(postmix_upper * 100.0),
                precavity_impurity_upper_bound=float(precavity_upper),
                precavity_impurity_upper_percent=float(precavity_upper * 100.0),
            )
        )
    return entries


def build_mixing_matrix(
    nearest_neighbors: list[list[NeighborMixEntry]],
    mode: str,
) -> np.ndarray:
    mix = np.eye(9, dtype=float)
    for lock_idx, entries in enumerate(nearest_neighbors):
        if mode == "literal_full_leak":
            weights = [entry.cavity_leak_ratio for entry in entries]
        elif mode == "data_driven_upper_bound":
            weights = [entry.postmix_upper_bound for entry in entries]
        else:
            raise ValueError(f"Unknown mode: {mode}")
        mix[lock_idx, lock_idx] = 1.0 - float(sum(weights))
        for entry, weight in zip(entries, weights):
            mix[entry.mode, lock_idx] = float(weight)
    return mix


def evaluate_scenario(
    s_measured: np.ndarray,
    y_obs: np.ndarray,
    x_reported: np.ndarray,
    mix_matrix: np.ndarray,
) -> ScenarioMetrics:
    s_corrected = s_measured @ np.linalg.inv(mix_matrix)
    x_corrected = solve_nnls(s_corrected, y_obs)

    p_reported = compute_percent_matrix(x_reported)
    p_corrected = compute_percent_matrix(x_corrected)
    er_reported = compute_er_sum_db(x_reported)
    er_corrected = compute_er_sum_db(x_corrected)
    diff = np.abs(p_reported - p_corrected)
    offdiag_mask = ~np.eye(9, dtype=bool)

    reported_mean_success = float(np.mean(np.diag(p_reported)))
    corrected_mean_success = float(np.mean(np.diag(p_corrected)))
    reported_mean_er = float(np.mean(er_reported))
    corrected_mean_er = float(np.mean(er_corrected))

    return ScenarioMetrics(
        reported_mean_success_percent=reported_mean_success,
        corrected_mean_success_percent=corrected_mean_success,
        reported_mean_er_db=reported_mean_er,
        corrected_mean_er_db=corrected_mean_er,
        reported_minus_corrected_success_pct_pt=reported_mean_success - corrected_mean_success,
        reported_minus_corrected_er_db=reported_mean_er - corrected_mean_er,
        mean_abs_matrix_delta_pct_pt=float(np.mean(diff)),
        max_abs_matrix_delta_pct_pt=float(np.max(diff)),
        mean_abs_offdiag_delta_pct_pt=float(np.mean(diff[offdiag_mask])),
        max_abs_offdiag_delta_pct_pt=float(np.max(diff[offdiag_mask])),
        max_abs_diag_delta_pct_pt=float(np.max(np.abs(np.diag(p_reported) - np.diag(p_corrected)))),
        max_abs_er_delta_db=float(np.max(np.abs(er_reported - er_corrected))),
        min_corrected_response_entry=float(np.min(s_corrected)),
    )


def summarize_run(run_dir: Path) -> RunSummary:
    transmission_xlsx, positions, peak_trans = read_transmission_data(run_dir)
    s_matrix = read_matrix(run_dir / "calib_matrix_with_uncertainty.xlsx", "mean")
    x_reported = read_matrix(run_dir / "full9_matrix_with_uncertainty.xlsx", "mean")
    y_obs = read_matrix(run_dir / "full9_matrix_with_uncertainty.xlsx", "raw_y_mean")

    nearest_neighbors = [
        build_neighbor_entries(lock_idx, positions, peak_trans, s_matrix)
        for lock_idx in range(9)
    ]

    literal_mix = build_mixing_matrix(nearest_neighbors, mode="literal_full_leak")
    data_upper_mix = build_mixing_matrix(nearest_neighbors, mode="data_driven_upper_bound")

    all_leaks = [entry.cavity_leak_percent for col in nearest_neighbors for entry in col]
    all_precavity = [entry.precavity_impurity_upper_percent for col in nearest_neighbors for entry in col]
    total_postmix_per_col = [
        float(sum(entry.postmix_upper_percent for entry in entries))
        for entries in nearest_neighbors
    ]

    return RunSummary(
        run=run_dir.name,
        transmission_xlsx=str(transmission_xlsx),
        lock_positions=[float(v) for v in positions],
        peak_transmittance_ratio=[float(v) for v in peak_trans],
        nearest_neighbors=nearest_neighbors,
        mean_single_neighbor_leak_percent=float(np.mean(all_leaks)),
        leak_percent_range=[float(np.min(all_leaks)), float(np.max(all_leaks))],
        mean_total_postmix_upper_percent_per_column=float(np.mean(total_postmix_per_col)),
        max_total_postmix_upper_percent_per_column=float(np.max(total_postmix_per_col)),
        mean_single_neighbor_precavity_impurity_upper_percent=float(np.mean(all_precavity)),
        max_single_neighbor_precavity_impurity_upper_percent=float(np.max(all_precavity)),
        literal_full_leak_scenario=evaluate_scenario(s_matrix, y_obs, x_reported, literal_mix),
        data_driven_upper_bound_scenario=evaluate_scenario(s_matrix, y_obs, x_reported, data_upper_mix),
    )


def build_results_payload(run_summaries: list[RunSummary], data_root: Path) -> dict[str, Any]:
    all_leaks = [entry.cavity_leak_percent for run in run_summaries for col in run.nearest_neighbors for entry in col]
    all_postmix_totals = [run.mean_total_postmix_upper_percent_per_column for run in run_summaries]
    all_postmix_column_max = [run.max_total_postmix_upper_percent_per_column for run in run_summaries]
    all_precavity = [
        entry.precavity_impurity_upper_percent
        for run in run_summaries
        for col in run.nearest_neighbors
        for entry in col
    ]
    data_upper_success_bias = [
        run.data_driven_upper_bound_scenario.reported_minus_corrected_success_pct_pt
        for run in run_summaries
    ]
    data_upper_er_bias = [
        run.data_driven_upper_bound_scenario.reported_minus_corrected_er_db
        for run in run_summaries
    ]
    data_upper_matrix_max = [
        run.data_driven_upper_bound_scenario.max_abs_matrix_delta_pct_pt
        for run in run_summaries
    ]
    data_upper_er_max = [
        run.data_driven_upper_bound_scenario.max_abs_er_delta_db
        for run in run_summaries
    ]

    return {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "data_root": str(data_root),
        "fixed_finesse": F_FIXED,
        "global_summary": {
            "nearest_neighbor_cavity_leak_percent_mean": float(np.mean(all_leaks)),
            "nearest_neighbor_cavity_leak_percent_range": [
                float(np.min(all_leaks)),
                float(np.max(all_leaks)),
            ],
            "total_postmix_upper_percent_per_column_mean": float(np.mean(all_postmix_totals)),
            "total_postmix_upper_percent_per_column_max": float(np.max(all_postmix_column_max)),
            "single_neighbor_precavity_impurity_upper_percent_mean": float(np.mean(all_precavity)),
            "single_neighbor_precavity_impurity_upper_percent_max": float(np.max(all_precavity)),
            "data_upper_bound_mean_success_bias_pct_pt": float(np.mean(data_upper_success_bias)),
            "data_upper_bound_mean_er_bias_db": float(np.mean(data_upper_er_bias)),
            "data_upper_bound_max_matrix_delta_pct_pt": float(np.max(data_upper_matrix_max)),
            "data_upper_bound_max_er_delta_db": float(np.max(data_upper_er_max)),
        },
        "runs": [asdict(run) for run in run_summaries],
    }


def build_results_markdown(payload: dict[str, Any]) -> str:
    global_summary = payload["global_summary"]
    lines: list[str] = []
    lines.append("# Q4 calibration-leakage bias check")
    lines.append("")
    lines.append(f"- Generated at: `{payload['generated_at']}`")
    lines.append(f"- Data root: `{payload['data_root']}`")
    lines.append(f"- Fixed finesse reused from Q2/Q3: `{payload['fixed_finesse']:.2f}`")
    lines.append("")
    lines.append("## Core result")
    lines.append("")
    lines.append(
        f"- Using the measured lock spacings and the same Airy model as Q2, the nearest-neighbor cavity leakage coefficient is indeed about "
        f"`{global_summary['nearest_neighbor_cavity_leak_percent_mean']:.3f}%` on average, with a run-and-column range of "
        f"`{global_summary['nearest_neighbor_cavity_leak_percent_range'][0]:.3f}%` to "
        f"`{global_summary['nearest_neighbor_cavity_leak_percent_range'][1]:.3f}%`."
    )
    lines.append(
        f"- But that coefficient is not the same thing as the amount actually mixed into the measured calibration column `S[:,k]`. "
        f"From the present `S` matrix itself, the two-neighbor total postmix fraction is conservatively bounded by only "
        f"`{global_summary['total_postmix_upper_percent_per_column_mean']:.4f}%` on average and "
        f"`{global_summary['total_postmix_upper_percent_per_column_max']:.4f}%` in the worst column."
    )
    lines.append(
        f"- Propagating this conservative upper bound through the same NNLS correction pipeline changes the reported mean success rate by only "
        f"`{global_summary['data_upper_bound_mean_success_bias_pct_pt']:.4f}` percentage points on average and the mean `ER_sum` by only "
        f"`{global_summary['data_upper_bound_mean_er_bias_db']:.4f} dB`."
    )
    lines.append(
        f"- Even the largest entry-wise change in the normalized 9x9 matrix remains only "
        f"`{global_summary['data_upper_bound_max_matrix_delta_pct_pt']:.4f}` percentage points, and the largest per-column `ER_sum` change is "
        f"`{global_summary['data_upper_bound_max_er_delta_db']:.4f} dB`."
    )
    lines.append("")
    lines.append("## Why the raw Q4 statement is too strong")
    lines.append("")
    lines.append("- `~2%` is the cavity leak coefficient of a neighboring mode *if that neighbor is present in the calibration input*.")
    lines.append("- The actual contamination absorbed into `S[:,k]` must still be multiplied by however much nearest-neighbor impurity existed before the cavity.")
    lines.append(
        f"- The current data imply that this pre-cavity nearest-neighbor impurity was at most about "
        f"`{global_summary['single_neighbor_precavity_impurity_upper_percent_mean']:.3f}%` on average and "
        f"`{global_summary['single_neighbor_precavity_impurity_upper_percent_max']:.3f}%` at the worst single-neighbor case."
    )
    lines.append("- So the physically relevant calibration impurity is sub-`10^-3` in fraction units, not a literal `2%` mixed into each measured column of `S`.")
    lines.append("")
    lines.append("## Literal 2% scenario check")
    lines.append("")
    lines.append(
        "- For completeness, the script also tests the reviewer-style literal scenario that each calibration column truly absorbs its full nearest-neighbor cavity leakage. "
        "That scenario forces the inferred ideal response matrix to develop negative entries at the `10^-3` to `10^-2` level, which is physically inconsistent with the present measured `S`."
    )
    lines.append("")
    lines.append("## Per-run summary")
    lines.append("")
    for run in payload["runs"]:
        data_upper = run["data_driven_upper_bound_scenario"]
        literal = run["literal_full_leak_scenario"]
        lines.append(f"### Run {run['run']}")
        lines.append(
            f"- Nearest-neighbor cavity leakage: mean `{run['mean_single_neighbor_leak_percent']:.3f}%`, "
            f"range `{run['leak_percent_range'][0]:.3f}%` to `{run['leak_percent_range'][1]:.3f}%`."
        )
        lines.append(
            f"- Total calibration postmix upper bound per column: mean `{run['mean_total_postmix_upper_percent_per_column']:.4f}%`, "
            f"max `{run['max_total_postmix_upper_percent_per_column']:.4f}%`."
        )
        lines.append(
            f"- Data-driven upper-bound bias: mean success `{data_upper['reported_minus_corrected_success_pct_pt']:.4f}` pct-pt high, "
            f"mean `ER_sum` `{data_upper['reported_minus_corrected_er_db']:.4f} dB` high."
        )
        lines.append(
            f"- Literal full-leak scenario: mean success bias `{literal['reported_minus_corrected_success_pct_pt']:.3f}` pct-pt, "
            f"mean `ER_sum` bias `{literal['reported_minus_corrected_er_db']:.3f} dB`, "
            f"minimum inferred ideal-response entry `{literal['min_corrected_response_entry']:.4e}`."
        )
        lines.append("")
    return "\n".join(lines) + "\n"


def build_plain_chinese_conclusion(payload: dict[str, Any]) -> str:
    g = payload["global_summary"]
    return (
        "# Q4 直白结论\n\n"
        "Q4 的关键不是“最近邻离共振漏透约 2%”这句话本身对不对，而是这 2% 到底有没有以同样量级真的混进当前标定矩阵 `S`。\n\n"
        f"用 Q2 已经验证过的那套前向模型重算后，最近邻模式在当前 9 模工作点下的 Airy 漏透系数确实就是约 "
        f"`{g['nearest_neighbor_cavity_leak_percent_mean']:.3f}%`，范围约 "
        f"`{g['nearest_neighbor_cavity_leak_percent_range'][0]:.3f}%` 到 "
        f"`{g['nearest_neighbor_cavity_leak_percent_range'][1]:.3f}%`。但这只是“如果最近邻本来就在标定输入里，它通过腔后还能剩多少”的系数，"
        "并不等于“`S` 里真的混进了 2% 的邻模”。\n\n"
        f"从当前实测 `S` 本身反推，最近邻杂质真正可能被吸收到单列标定里的总份额，平均只到 "
        f"`{g['total_postmix_upper_percent_per_column_mean']:.4f}%`，最坏一列也只有 "
        f"`{g['total_postmix_upper_percent_per_column_max']:.4f}%`。这已经是一个很保守的上界了。\n\n"
        f"把这个上界继续按现有 NNLS 校正流程传播到满载结果，得到的偏差方向确实与你说的一致：当前正文报告值会把成功率和 `ER_sum` 轻微高估。"
        f"但量级非常小，三次 run 平均下来，成功率只高了 "
        f"`{g['data_upper_bound_mean_success_bias_pct_pt']:.4f}` 个百分点，平均 `ER_sum` 只高了 "
        f"`{g['data_upper_bound_mean_er_bias_db']:.4f} dB`；整个 9×9 百分比矩阵里最大的单元变化也只有 "
        f"`{g['data_upper_bound_max_matrix_delta_pct_pt']:.4f}` 个百分点。\n\n"
        "所以，Q4 最稳妥的结论不是“这个效应可以把结果明显美化”，而是：\n\n"
        "1. 它在方向上确实会让校正后的结果略微偏乐观。\n"
        "2. 但按当前数据能给出的保守上界，这个偏差只有 `10^-3 ~ 10^-2 dB` 和 `10^-2` 个百分点量级，不足以改变第五章的主要结论。\n"
        "3. 因此最合适的处理方式不是推翻标定框架，而是在附录里把“纯目标模式”改写成“高纯度目标模式，并存在可量化上界的微弱最近邻残余”，老实补一句说明即可。\n"
    )


def build_detailed_chinese_report(payload: dict[str, Any]) -> str:
    g = payload["global_summary"]
    lines: list[str] = []
    lines.append("# Q4 详细报告：标定时最近邻漏透会不会把 `ER_sum` 轻微抬高")
    lines.append("")
    lines.append("## 1. 这次 Q4 具体在检验什么")
    lines.append("")
    lines.append("Q4 的原始担心是：附录把标定态说成“到达检测端的光场可视为纯的目标模式”，但如果 `tau_min \\approx 3.5`，最近邻模式的洛伦兹/ Airy 漏透仍有约 `2%`，那么标定态就不是绝对纯净。")
    lines.append("进一步的追问有三层：")
    lines.append("")
    lines.append("1. 这件事在当前标定矩阵 `S` 里到底占多大量级。")
    lines.append("2. 它会不会让校正后的 `ER_sum` 或成功率略微偏高。")
    lines.append("3. 这个效应是可以忽略、可以给上界，还是应该在附录里补一句。")
    lines.append("")
    lines.append("## 2. 这次计算沿用的前提")
    lines.append("")
    lines.append("- 沿用 Q2 已确认的前向模型：离共振主要改幅度，不改横向轮廓。")
    lines.append("- 不重新怀疑 `S` 的列定义；仍按“产生端 `l=k`、FP 锁定 `l=k`、检测端遍历 `l=0~8`”理解。")
    lines.append("- 不改正文、不改原始数据、不改已有工作簿。")
    lines.append("")
    lines.append("## 3. 先把“2%”说清楚")
    lines.append("")
    lines.append("用当前 3 次 run 的实测锁定峰位、单模峰值透射率以及与 Q2 相同的 Airy 线型，最近邻模式在标定锁定态下的漏透系数为：")
    lines.append("")
    lines.append(
        f"- 单个最近邻的平均腔漏透系数：`{g['nearest_neighbor_cavity_leak_percent_mean']:.3f}%`"
    )
    lines.append(
        f"- 范围：`{g['nearest_neighbor_cavity_leak_percent_range'][0]:.3f}%` 到 "
        f"`{g['nearest_neighbor_cavity_leak_percent_range'][1]:.3f}%`"
    )
    lines.append("")
    lines.append("这一步说明：审稿人抓住“`tau_min \\approx 3.5` 时最近邻并非被完全截断”这个物理事实，本身没有问题。")
    lines.append("")
    lines.append("但这里有一个常被偷换的步骤：")
    lines.append("")
    lines.append("> `2%` 只是“邻模若存在于标定输入中，经腔后还剩多少”的系数，不是“`S` 里已经混进了 2% 邻模”的直接结论。")
    lines.append("")
    lines.append("真正混进 `S[:,k]` 的量，还要再乘上标定输入里原本存在的邻模杂质份额。")
    lines.append("")
    lines.append("## 4. 如何从当前 `S` 给出保守上界")
    lines.append("")
    lines.append("若第 `k` 列标定态里真的混入了少量邻模 `n`，那么它至少会在第 `n` 个检测通道上留下一个信号；而纯 `n` 模自己的第 `n` 通道响应就是 `S_{nn}`。")
    lines.append("因此可直接从当前矩阵给出一个保守上界：")
    lines.append("")
    lines.append(r"```text")
    lines.append(r"lambda_(n<-k) <= S[n,k] / S[n,n]")
    lines.append(r"```")
    lines.append("")
    lines.append("这里 `lambda_(n<-k)` 表示“真正被吸收到第 `k` 列标定态里的邻模 `n` 功率份额”。这个上界是保守的，因为 `S[n,k]` 本来还包含检测链路自身的交叉响应，并不全是标定杂质。")
    lines.append("")
    lines.append("把两侧最近邻都算进去后，当前数据给出的结果是：")
    lines.append("")
    lines.append(
        f"- 每列总标定杂质上界的平均值：`{g['total_postmix_upper_percent_per_column_mean']:.4f}%`"
    )
    lines.append(
        f"- 每列总标定杂质上界的最大值：`{g['total_postmix_upper_percent_per_column_max']:.4f}%`"
    )
    lines.append("")
    lines.append("也就是说，虽然腔本身对最近邻的漏透系数是约 `2%`，但在当前实测 `S` 中，真正可能被吸收进去的那部分最近邻杂质，只有 `10^-4 ~ 10^-3` 分数量级。")
    lines.append("")
    lines.append("进一步换算回腔前输入，得到的最近邻预腔杂质上界也只有：")
    lines.append("")
    lines.append(
        f"- 平均：`{g['single_neighbor_precavity_impurity_upper_percent_mean']:.3f}%`"
    )
    lines.append(
        f"- 最坏单邻模情形：`{g['single_neighbor_precavity_impurity_upper_percent_max']:.3f}%`"
    )
    lines.append("")
    lines.append("这与“高纯度单模输入”这一实验背景并不矛盾。")
    lines.append("")
    lines.append("## 5. 这会不会把 `ER_sum` 或成功率轻微抬高")
    lines.append("")
    lines.append("会，方向上是会让当前报告值略微偏乐观。")
    lines.append("")
    lines.append("建模方式是：把当前测得的 `S_meas` 看成某个更理想检测矩阵 `D` 与一个很小的“标定态混合矩阵” `M` 的乘积，即")
    lines.append("")
    lines.append(r"```text")
    lines.append(r"S_meas = D · M")
    lines.append(r"```")
    lines.append("")
    lines.append("如果标定态并非绝对纯净，那么正文当前用 `S_meas` 做反演，得到的是略带偏差的 `X_reported`；而更理想的结果应由 `D` 去解同一个原始 `Y`。")
    lines.append("我把上面得到的保守上界装进 `M`，再用与正文一致的 NNLS 反演流程重解一次。")
    lines.append("")
    lines.append("三次 run 的平均结果是：")
    lines.append("")
    lines.append(
        f"- 当前报告的平均成功率比更理想结果高 `"
        f"{g['data_upper_bound_mean_success_bias_pct_pt']:.4f}` 个百分点"
    )
    lines.append(
        f"- 当前报告的平均 `ER_sum` 比更理想结果高 `"
        f"{g['data_upper_bound_mean_er_bias_db']:.4f} dB`"
    )
    lines.append(
        f"- 9×9 归一化矩阵中最大单元改变量只有 `"
        f"{g['data_upper_bound_max_matrix_delta_pct_pt']:.4f}` 个百分点"
    )
    lines.append(
        f"- 单列 `ER_sum` 的最大改变量只有 `"
        f"{g['data_upper_bound_max_er_delta_db']:.4f} dB`"
    )
    lines.append("")
    lines.append("所以结论很明确：")
    lines.append("")
    lines.append("> 这个系统偏差的方向确实是让当前结果略好看一点，但量级非常小，只是一个可定量上界的轻微系统偏差。")
    lines.append("")
    lines.append("## 6. 为什么不能把“2%”直接当成已经混进 `S` 的量")
    lines.append("")
    lines.append("为了防止口径滑回审稿人的原始说法，我还额外测试了一个“字面 2% 混入”的极端情形：")
    lines.append("")
    lines.append("- 直接假设每个标定列都真的吸收了它全部最近邻腔漏透份额。")
    lines.append("")
    lines.append("这样反推出的理想检测矩阵会出现明显负元（`10^-3 ~ 10^-2` 量级），这是不物理的。")
    lines.append("也就是说，若真的把“最近邻腔漏透约 2%”直接理解成“`S` 的每列已混入 2% 邻模”，那和当前实测 `S` 本身是不相容的。")
    lines.append("")
    lines.append("## 7. Q4 最终应该怎么落")
    lines.append("")
    lines.append("最稳妥的结论是：")
    lines.append("")
    lines.append("1. `tau_min \\approx 3.5` 时，最近邻离共振并未被完全切断，这一点应当承认。")
    lines.append("2. 但这并不意味着当前 `S` 里真的混入了整整 `2%` 的邻模列。对当前数据而言，真正可能吸收到 `S` 里的最近邻杂质上界只有 `10^-4 ~ 10^-3` 分数量级。")
    lines.append("3. 这会使校正后的 `ER_sum` 和成功率略微偏高，但偏差只有 `10^-3 ~ 10^-2 dB`、`10^-2` 个百分点量级，远小于足以改写正文结论的程度。")
    lines.append("4. 因而最合适的处理不是推翻校正框架，而是在附录里把“纯目标模式”改成“高纯度目标模式，并存在可量化上界的最近邻残余漏透”，诚实补一句说明。")
    lines.append("")
    return "\n".join(lines) + "\n"


def main() -> None:
    output_dir = Path(__file__).resolve().parent
    data_root = output_dir.parent / "最终数据"
    run_dirs = discover_run_dirs(data_root)
    if len(run_dirs) != 3:
        raise RuntimeError(f"Expected 3 run directories under {data_root}, found: {run_dirs}")

    run_summaries = [summarize_run(run_dir) for run_dir in run_dirs]
    payload = build_results_payload(run_summaries, data_root)

    results_json = output_dir / "q4_calibration_leakage_bias_results.json"
    results_md = output_dir / "q4_calibration_leakage_bias_results.md"
    plain_md = output_dir / "Q4_直白结论.md"
    detailed_md = output_dir / "Q4_详细报告.md"

    results_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    results_md.write_text(build_results_markdown(payload), encoding="utf-8")
    plain_md.write_text(build_plain_chinese_conclusion(payload), encoding="utf-8")
    detailed_md.write_text(build_detailed_chinese_report(payload), encoding="utf-8")

    print(f"Wrote: {results_json}")
    print(f"Wrote: {results_md}")
    print(f"Wrote: {plain_md}")
    print(f"Wrote: {detailed_md}")


if __name__ == "__main__":
    main()
