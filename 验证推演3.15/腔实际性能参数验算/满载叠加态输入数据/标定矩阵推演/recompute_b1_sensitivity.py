#!/usr/bin/env python3
"""Recompute Appendix B1 sensitivity analysis from trusted triplicate data.

This script is intentionally standalone and read-only with respect to the
thesis source. It loads the trusted per-run calibration/full-load matrices from
the sibling data directory, recomputes the de-embedded ER_sum baseline, applies
several impurity perturbation models to the calibration matrix, and writes a
plain-text report for audit.
"""

from __future__ import annotations

import argparse
import math
import re
from dataclasses import dataclass
from pathlib import Path
from statistics import mean
from typing import Iterable

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import nnls


LOCK_RE = re.compile(r"lock\s*=\s*(\d+)", re.IGNORECASE)
CH_RE = re.compile(r"ch\s*=\s*(\d+)", re.IGNORECASE)
APP_ROW_RE = re.compile(
    r"^\s*(\d+)\s*&\s*([0-9.]+)\s*&\s*([0-9.]+)\s*&\s*([0-9.]+)\s*\\\\"
)
LEGACY_ROW_RE = re.compile(
    r"^\s*(\d+)\s*&\s*([0-9.]+)\s*&\s*([0-9.]+)\s*&\s*([0-9.]+)\s*&\s*(\d+)\s*&"
)


@dataclass
class MatrixAxes:
    channels: list[int]
    locks: list[int]


@dataclass
class RunData:
    run_name: str
    axes: MatrixAxes
    s_matrix: np.ndarray
    y_matrix: np.ndarray


@dataclass
class RunResult:
    run_name: str
    axes: MatrixAxes
    x_matrix: np.ndarray
    pct_matrix: np.ndarray
    er_sum_db: dict[int, float]
    er_max_db: dict[int, float]
    signal_share: dict[int, float]


def parse_float(value) -> float:
    if value is None:
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text in {"", "NaN", "nan"}:
        return float("nan")
    if text == "Inf":
        return float("inf")
    if text == "-Inf":
        return float("-inf")
    return float(text)


def parse_matrix_sheet(path: Path, sheet_name: str) -> tuple[MatrixAxes, np.ndarray]:
    wb = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise KeyError(f"{path} missing sheet {sheet_name}")
    ws = wb[sheet_name]

    locks: list[int] = []
    for col in range(2, ws.max_column + 1):
        header = ws.cell(1, col).value
        if header is None:
            continue
        match = LOCK_RE.search(str(header))
        if match:
            locks.append(int(match.group(1)))

    channels: list[int] = []
    rows: list[list[float]] = []
    for row in range(2, ws.max_row + 1):
        label = ws.cell(row, 1).value
        if label is None:
            continue
        match = CH_RE.search(str(label))
        if not match:
            continue
        channels.append(int(match.group(1)))
        rows.append([parse_float(ws.cell(row, col).value) for col in range(2, 2 + len(locks))])

    wb.close()
    axes = MatrixAxes(channels=channels, locks=locks)
    return axes, np.asarray(rows, dtype=float)


def discover_data_dir(script_dir: Path) -> Path:
    parent = script_dir.parent
    candidates = [p for p in parent.iterdir() if p.is_dir() and (p / "triplicate_full9_aggregate.xlsx").exists()]
    if len(candidates) != 1:
        raise RuntimeError(f"Could not uniquely discover trusted data directory under {parent}")
    return candidates[0]


def load_runs(data_dir: Path, runs: Iterable[str]) -> list[RunData]:
    run_data: list[RunData] = []
    for run_name in runs:
        run_dir = data_dir / run_name
        calib_path = run_dir / "calib_matrix_with_uncertainty.xlsx"
        full9_path = run_dir / "full9_matrix_with_uncertainty.xlsx"
        axes_s, s_matrix = parse_matrix_sheet(calib_path, "mean")
        # `mean` in the full9 workbook is already the de-embedded x-matrix.
        # For a sensitivity re-solve we must start from the raw observation y.
        axes_y, y_matrix = parse_matrix_sheet(full9_path, "raw_y_mean")
        if axes_s != axes_y:
            raise ValueError(f"Axis mismatch in run {run_name}")
        run_data.append(RunData(run_name=run_name, axes=axes_s, s_matrix=s_matrix, y_matrix=y_matrix))
    return run_data


def safe_db10(num: float, den: float) -> float:
    if not math.isfinite(num) or not math.isfinite(den) or num <= 0 or den <= 0:
        return float("nan")
    return 10.0 * math.log10(num / den)


def solve_nnls_matrix(s_matrix: np.ndarray, y_matrix: np.ndarray) -> np.ndarray:
    x_matrix = np.zeros_like(y_matrix, dtype=float)
    for col in range(y_matrix.shape[1]):
        y = np.clip(np.nan_to_num(y_matrix[:, col], nan=0.0), 0.0, None)
        sol, _ = nnls(s_matrix, y)
        x_matrix[:, col] = sol
    return x_matrix


def normalize_columns(x_matrix: np.ndarray) -> np.ndarray:
    col_sums = np.sum(x_matrix, axis=0, keepdims=True)
    return np.divide(x_matrix, col_sums, out=np.zeros_like(x_matrix), where=col_sums > 0)


def compute_metrics(x_matrix: np.ndarray, axes: MatrixAxes) -> tuple[dict[int, float], dict[int, float], dict[int, float]]:
    ch_to_row = {ch: idx for idx, ch in enumerate(axes.channels)}
    er_sum_db: dict[int, float] = {}
    er_max_db: dict[int, float] = {}
    signal_share: dict[int, float] = {}

    for col, lock in enumerate(axes.locks):
        if lock not in ch_to_row:
            continue
        sig_row = ch_to_row[lock]
        col_vec = np.clip(x_matrix[:, col], 0.0, None)
        signal = float(col_vec[sig_row])
        leak = np.delete(col_vec, sig_row)
        leak_sum = float(np.sum(leak)) if leak.size else float("nan")
        leak_max = float(np.max(leak)) if leak.size else float("nan")
        total = float(np.sum(col_vec))
        er_sum_db[lock] = safe_db10(signal, leak_sum)
        er_max_db[lock] = safe_db10(signal, leak_max)
        signal_share[lock] = signal / total if total > 0 else float("nan")

    return er_sum_db, er_max_db, signal_share


def build_run_result(run: RunData, s_matrix: np.ndarray | None = None) -> RunResult:
    effective_s = run.s_matrix if s_matrix is None else s_matrix
    x_matrix = solve_nnls_matrix(effective_s, run.y_matrix)
    pct_matrix = normalize_columns(x_matrix)
    er_sum_db, er_max_db, signal_share = compute_metrics(x_matrix, run.axes)
    return RunResult(
        run_name=run.run_name,
        axes=run.axes,
        x_matrix=x_matrix,
        pct_matrix=pct_matrix,
        er_sum_db=er_sum_db,
        er_max_db=er_max_db,
        signal_share=signal_share,
    )


def aggregate_metric(results: list[RunResult], key: str) -> dict[int, tuple[float, float]]:
    locks = results[0].axes.locks
    out: dict[int, tuple[float, float]] = {}
    for lock in locks:
        vals = [getattr(result, key)[lock] for result in results]
        out[lock] = (float(np.mean(vals)), float(np.std(vals, ddof=1)) if len(vals) > 1 else 0.0)
    return out


def combined_pct_matrix(results: list[RunResult]) -> np.ndarray:
    return np.mean(np.stack([result.pct_matrix for result in results], axis=0), axis=0)


def make_single_mapping_from_pct(pct_matrix: np.ndarray, axes: MatrixAxes) -> dict[int, list[tuple[int, float]]]:
    mapping: dict[int, list[tuple[int, float]]] = {}
    ch_to_row = {ch: idx for idx, ch in enumerate(axes.channels)}
    for col, lock in enumerate(axes.locks):
        candidates = []
        for row, ch in enumerate(axes.channels):
            if ch == lock:
                continue
            candidates.append((pct_matrix[row, col], ch))
        _, best_ch = max(candidates, key=lambda item: item[0])
        mapping[lock] = [(best_ch, 1.0)]
    return mapping


def make_foldback_single_mapping(pct_matrix: np.ndarray, axes: MatrixAxes) -> dict[int, list[tuple[int, float]]]:
    mapping: dict[int, list[tuple[int, float]]] = {}
    ch_to_row = {ch: idx for idx, ch in enumerate(axes.channels)}
    lock_to_col = {lock: idx for idx, lock in enumerate(axes.locks)}
    for lock in axes.locks:
        candidates = [((lock + 4) % 9), ((lock + 5) % 9)]
        col = lock_to_col[lock]
        ranked = [(pct_matrix[ch_to_row[ch], col], ch) for ch in candidates]
        _, best_ch = max(ranked, key=lambda item: item[0])
        mapping[lock] = [(best_ch, 1.0)]
    return mapping


def make_foldback_split_mapping() -> dict[int, list[tuple[int, float]]]:
    mapping: dict[int, list[tuple[int, float]]] = {}
    for lock in range(9):
        mapping[lock] = [((lock + 4) % 9, 0.5), ((lock + 5) % 9, 0.5)]
    return mapping


def legacy_mapping() -> dict[int, list[tuple[int, float]]]:
    return {
        0: [(5, 1.0)],
        1: [(6, 1.0)],
        2: [(7, 1.0)],
        3: [(8, 1.0)],
        4: [(8, 1.0)],
        5: [(0, 1.0)],
        6: [(2, 1.0)],
        7: [(2, 1.0)],
        8: [(4, 1.0)],
    }


def perturb_s_matrix(run: RunData, mapping: dict[int, list[tuple[int, float]]], impurity_fraction: float) -> np.ndarray:
    if impurity_fraction < 0 or impurity_fraction >= 1:
        raise ValueError("impurity_fraction must lie in [0,1)")
    s_new = run.s_matrix.copy()
    lock_to_col = {lock: idx for idx, lock in enumerate(run.axes.locks)}
    for lock in run.axes.locks:
        col = lock_to_col[lock]
        sources = mapping[lock]
        if not sources:
            continue
        mixed = (1.0 - impurity_fraction) * run.s_matrix[:, col]
        for source_lock, weight in sources:
            src_col = lock_to_col[source_lock]
            mixed += impurity_fraction * weight * run.s_matrix[:, src_col]
        s_new[:, col] = mixed
    return s_new


def extract_current_appendix_table(path: Path) -> dict[int, tuple[float, float, float]]:
    rows: dict[int, tuple[float, float, float]] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        match = APP_ROW_RE.match(line)
        if match:
            lock = int(match.group(1))
            rows[lock] = (float(match.group(2)), float(match.group(3)), float(match.group(4)))
    return rows


def extract_legacy_table(path: Path) -> dict[int, tuple[float, float]]:
    rows: dict[int, tuple[float, float]] = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        match = LEGACY_ROW_RE.match(line)
        if match:
            lock = int(match.group(1))
            er_max = float(match.group(2))
            er_sum = float(match.group(3))
            rows[lock] = (er_max, er_sum)
    return rows


def find_first_existing(paths: Iterable[Path]) -> Path | None:
    for path in paths:
        if path.exists():
            return path
    return None


def format_mapping(mapping: dict[int, list[tuple[int, float]]]) -> str:
    parts = []
    for lock in sorted(mapping):
        sources = ", ".join(f"{src}:{weight:.2f}" for src, weight in mapping[lock])
        parts.append(f"{lock}->{sources}")
    return "; ".join(parts)


def compare_vectors(a: dict[int, float], b: dict[int, float], tol: float = 1e-9) -> bool:
    if sorted(a) != sorted(b):
        return False
    for key in a:
        if abs(a[key] - b[key]) > tol:
            return False
    return True


def render_metric_table(
    title: str,
    baseline: dict[int, tuple[float, float]],
    perturbed: dict[int, tuple[float, float]],
) -> list[str]:
    lines = [title, "l  baseline_mean  baseline_sd  perturbed_mean  perturbed_sd  delta_mean"]
    for lock in sorted(baseline):
        base_mean, base_sd = baseline[lock]
        pert_mean, pert_sd = perturbed[lock]
        lines.append(
            f"{lock:<2} {base_mean:>13.3f} {base_sd:>11.3f} {pert_mean:>15.3f} {pert_sd:>12.3f} {pert_mean - base_mean:>11.3f}"
        )
    return lines


def main() -> None:
    script_dir = Path(__file__).resolve().parent
    default_data_dir = discover_data_dir(script_dir)
    parser = argparse.ArgumentParser(description="Recompute Appendix B1 sensitivity from trusted triplicate data.")
    parser.add_argument("--data-dir", type=Path, default=default_data_dir)
    parser.add_argument("--runs", nargs="+", default=["1", "2", "3"])
    parser.add_argument("--output", type=Path, default=script_dir / "b1_sensitivity_report.txt")
    parser.add_argument("--appendix-tex", type=Path, default=None)
    parser.add_argument("--legacy-tex", type=Path, default=None)
    args = parser.parse_args()

    run_data = load_runs(args.data_dir, args.runs)
    baseline_results = [build_run_result(run) for run in run_data]
    baseline_er = aggregate_metric(baseline_results, "er_sum_db")
    pct_combined = combined_pct_matrix(baseline_results)
    axes = baseline_results[0].axes

    mappings = {
        "combined_max_offdiag": make_single_mapping_from_pct(pct_combined, axes),
        "foldback_single": make_foldback_single_mapping(pct_combined, axes),
        "foldback_split": make_foldback_split_mapping(),
        "legacy_single": legacy_mapping(),
    }
    impurity_levels = {
        "2.00%": 0.0200,
        "2.44%": 0.0244,
    }

    lines: list[str] = []
    lines.append("Appendix B1 trusted-data recomputation report")
    lines.append(f"data_dir = {args.data_dir}")
    lines.append(f"runs = {', '.join(args.runs)}")
    lines.append("")
    lines.append("Trusted baseline ER_sum from current triplicate data (mean +/- sd across runs):")
    for lock in axes.locks:
        base_mean, base_sd = baseline_er[lock]
        lines.append(f"  l={lock}: {base_mean:.3f} +/- {base_sd:.3f} dB")
    lines.append(f"  overall mean = {mean([baseline_er[lock][0] for lock in axes.locks]):.3f} dB")
    lines.append("")

    appendix_path = args.appendix_tex
    legacy_path = args.legacy_tex
    if appendix_path is None:
        appendix_path = find_first_existing(
            [
                Path.cwd() / "appendices" / "app_e5_calibration_robustness.tex",
                Path.cwd().parent / "thesis" / "appendices" / "app_e5_calibration_robustness.tex",
                script_dir.parents[4] / "thesis" / "appendices" / "app_e5_calibration_robustness.tex",
            ]
        )
    if legacy_path is None:
        legacy_path = find_first_existing(
            [
                script_dir.parents[3]
                / "_AI_DO_NOT_READ_半成品隔离区"
                / "THESIS_THESIS_LEGACY"
                / "备份"
                / "ch05_fp_design_and_validation旧版.tex",
                script_dir.parents[3]
                / "_AI_DO_NOT_READ_半成品隔离区"
                / "THESIS_THESIS_LEGACY"
                / "备份"
                / "ch05_fp_design_and_validation.tex",
                script_dir.parents[4]
                / "_AI_DO_NOT_READ_半成品隔离区"
                / "THESIS_THESIS_LEGACY"
                / "备份"
                / "ch05_fp_design_and_validation旧版.tex",
                script_dir.parents[4]
                / "_AI_DO_NOT_READ_半成品隔离区"
                / "THESIS_THESIS_LEGACY"
                / "备份"
                / "ch05_fp_design_and_validation.tex",
            ]
        )

    if appendix_path and appendix_path.exists():
        appendix_rows = extract_current_appendix_table(appendix_path)
        appendix_baseline = {lock: appendix_rows[lock][0] for lock in sorted(appendix_rows)}
        lines.append(f"Current appendix table detected: {appendix_path}")
        for lock in sorted(appendix_baseline):
            lines.append(f"  appendix l={lock}: baseline={appendix_baseline[lock]:.2f} dB")
        lines.append("")
    else:
        appendix_rows = {}
        appendix_baseline = {}

    if legacy_path and legacy_path.exists():
        legacy_rows = extract_legacy_table(legacy_path)
        legacy_er_sum = {lock: legacy_rows[lock][1] for lock in sorted(legacy_rows)}
        lines.append(f"Legacy full-load table detected: {legacy_path}")
        exact_match = compare_vectors(appendix_baseline, legacy_er_sum) if appendix_baseline else False
        lines.append(f"  appendix-baseline == legacy ER_sum exact match: {exact_match}")
        if exact_match:
            lines.append("  conclusion: the current appendix baseline column is copied from the legacy chapter-5 ER_sum table.")
        lines.append("")
    else:
        legacy_rows = {}

    trusted_baseline_rounded = {lock: round(baseline_er[lock][0], 2) for lock in axes.locks}
    if appendix_baseline:
        diffs = {lock: appendix_baseline[lock] - baseline_er[lock][0] for lock in axes.locks}
        lines.append("Difference between current appendix baseline and trusted triplicate baseline:")
        for lock in axes.locks:
            lines.append(
                f"  l={lock}: appendix={appendix_baseline[lock]:.2f} dB, trusted={baseline_er[lock][0]:.3f} dB, delta={diffs[lock]:+.3f} dB"
            )
        lines.append("")
        lines.append("Trusted baseline rounded to 0.01 dB (recommended if Appendix B1 is kept):")
        lines.append("  " + ", ".join(f"l={lock}:{trusted_baseline_rounded[lock]:.2f}" for lock in axes.locks))
        lines.append("")

    for impurity_label, impurity_fraction in impurity_levels.items():
        lines.append(f"Impurity level = {impurity_label} of total calibration-column weight")
        for mapping_name, mapping in mappings.items():
            perturbed_results = [
                build_run_result(run, perturb_s_matrix(run, mapping, impurity_fraction))
                for run in run_data
            ]
            perturbed_er = aggregate_metric(perturbed_results, "er_sum_db")
            lines.append(f"  mapping = {mapping_name}")
            lines.append(f"    source map: {format_mapping(mapping)}")
            delta_vals = [perturbed_er[lock][0] - baseline_er[lock][0] for lock in axes.locks]
            lines.append(
                "    delta summary: "
                f"mean={float(np.mean(delta_vals)):+.3f} dB, "
                f"min={float(np.min(delta_vals)):+.3f} dB, "
                f"max={float(np.max(delta_vals)):+.3f} dB"
            )
            for lock in axes.locks:
                base_mean, base_sd = baseline_er[lock]
                pert_mean, pert_sd = perturbed_er[lock]
                lines.append(
                    f"    l={lock}: baseline={base_mean:.3f}+/-{base_sd:.3f} dB, "
                    f"perturbed={pert_mean:.3f}+/-{pert_sd:.3f} dB, "
                    f"delta={pert_mean - base_mean:+.3f} dB"
                )
            lines.append("")

    # Recommended table: current thesis logic is closest to the foldback-single, 2% model.
    recommended_mapping = mappings["foldback_single"]
    recommended_results = [
        build_run_result(run, perturb_s_matrix(run, recommended_mapping, impurity_levels["2.00%"]))
        for run in run_data
    ]
    recommended_er = aggregate_metric(recommended_results, "er_sum_db")
    lines.extend(
        render_metric_table(
            "Recommended replacement for Appendix B1 under current trusted-data logic (2.00%, foldback_single):",
            baseline_er,
            recommended_er,
        )
    )
    lines.append("")

    args.output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print("\n".join(lines))
    print(f"\nSaved report to: {args.output}")


if __name__ == "__main__":
    main()
