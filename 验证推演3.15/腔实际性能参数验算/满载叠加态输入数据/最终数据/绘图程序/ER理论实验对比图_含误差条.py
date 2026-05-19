"""Plot ER_sum / ER_pairmax theory-vs-experiment charts with error bars.

Experimental data source:
  - triplicate_full9_aggregate.xlsx (sheet: ER_summary)

Theory data source:
  - tao_snr_l0_to_l8.csv
    - ER_full_actual_dB   -> for ER_sum comparison
    - C_pair_nn_dB        -> for ER_pairmax comparison

Outputs:
  - 输出图/ER_sum理论实验对比_含误差条.pdf
  - 输出图/ER_pairmax理论实验对比_含误差条.pdf
"""

from __future__ import annotations

import argparse
import csv
import math
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np
from openpyxl import load_workbook


T975_DF2 = 4.302652729911275  # t(0.975, df=2), n=3
ERRORBAR_MODES = ("ci95_t", "sd", "sem", "minmax", "scatter", "scatter_sd")
MODE_TO_SUFFIX = {
    "ci95_t": "CI95_t",
    "sd": "SD",
    "sem": "SEM",
    "minmax": "MinMax",
    "scatter": "Scatter",
    "scatter_sd": "Scatter_SD",
}
MODE_TO_LABEL = {
    "ci95_t": r"95\%CI (t, $n=3$)",
    "sd": r"SD ($n=3$)",
    "sem": r"SEM ($n=3$)",
    "minmax": r"Min-Max envelope ($n=3$)",
    "scatter": r"Raw data ($n=3$)",
    "scatter_sd": r"Raw data + SD",
}


def configure_plot_style() -> None:
    mpl.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["STIX Two Text", "STIXGeneral", "Times New Roman"],
            "mathtext.fontset": "stix",
            "font.size": 9,
            "axes.labelsize": 10,
            "axes.linewidth": 0.6,
            "axes.grid": False,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "xtick.direction": "in",
            "ytick.direction": "in",
            "xtick.major.width": 0.5,
            "ytick.major.width": 0.5,
            "xtick.major.size": 3.0,
            "ytick.major.size": 3.0,
            "legend.fontsize": 8,
            "legend.frameon": False,
            "grid.linewidth": 0.4,
            "grid.alpha": 0.5,
            "figure.dpi": 150,
        }
    )


def parse_er_summary(agg_xlsx: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    wb = load_workbook(agg_xlsx, read_only=True, data_only=True)
    if "ER_summary" not in wb.sheetnames:
        wb.close()
        raise KeyError(f"Missing sheet 'ER_summary' in {agg_xlsx.name}")

    ws = wb["ER_summary"]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    idx = {str(v): i + 1 for i, v in enumerate(header) if v is not None}

    required = [
        "channel(lock=l)",
        "ER_sum_run1_dB",
        "ER_sum_run2_dB",
        "ER_sum_run3_dB",
        "ER_max_run1_dB",
        "ER_max_run2_dB",
        "ER_max_run3_dB",
    ]
    for key in required:
        if key not in idx:
            wb.close()
            raise KeyError(f"Missing column '{key}' in ER_summary")

    channels: list[int] = []
    er_sum_runs: list[list[float]] = []
    er_max_runs: list[list[float]] = []

    for r in range(2, ws.max_row + 1):
        cv = ws.cell(r, idx["channel(lock=l)"]).value
        if cv is None:
            break
        ch = int(float(cv))
        channels.append(ch)

        er_sum_runs.append(
            [
                float(ws.cell(r, idx["ER_sum_run1_dB"]).value),
                float(ws.cell(r, idx["ER_sum_run2_dB"]).value),
                float(ws.cell(r, idx["ER_sum_run3_dB"]).value),
            ]
        )
        er_max_runs.append(
            [
                float(ws.cell(r, idx["ER_max_run1_dB"]).value),
                float(ws.cell(r, idx["ER_max_run2_dB"]).value),
                float(ws.cell(r, idx["ER_max_run3_dB"]).value),
            ]
        )

    wb.close()

    order = np.argsort(np.asarray(channels, dtype=int))
    ch_arr = np.asarray(channels, dtype=int)[order]
    er_sum = np.asarray(er_sum_runs, dtype=float)[order, :]
    er_max = np.asarray(er_max_runs, dtype=float)[order, :]
    return ch_arr, er_sum, er_max


def compute_error_stats(run_values: np.ndarray) -> dict[str, np.ndarray]:
    # run_values shape: (channels, 3)
    n = run_values.shape[1]
    mean = np.mean(run_values, axis=1)
    sd = np.std(run_values, axis=1, ddof=1)
    sem = sd / math.sqrt(n)
    ci95_t = T975_DF2 * sem
    min_v = np.min(run_values, axis=1)
    max_v = np.max(run_values, axis=1)
    return {
        "mean": mean,
        "sd": sd,
        "sem": sem,
        "ci95_t": ci95_t,
        "min": min_v,
        "max": max_v,
        "raw": run_values,
    }


def get_yerr_for_mode(stats: dict[str, np.ndarray], mode: str) -> np.ndarray:
    if mode == "minmax":
        lower = stats["mean"] - stats["min"]
        upper = stats["max"] - stats["mean"]
        return np.vstack([lower, upper])
    return stats[mode]


def parse_theory_csv(theory_csv: Path, channels: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    theory_sum_map: dict[int, float] = {}
    theory_pair_map: dict[int, float] = {}

    with theory_csv.open("r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        required = {"l", "ER_full_actual_dB", "C_pair_nn_dB"}
        if not reader.fieldnames or not required.issubset(set(reader.fieldnames)):
            raise KeyError(
                f"Theory CSV missing required columns: {sorted(required)}; got {reader.fieldnames}"
            )

        for row in reader:
            l = int(float(row["l"]))
            theory_sum_map[l] = float(row["ER_full_actual_dB"])
            theory_pair_map[l] = float(row["C_pair_nn_dB"])

    theory_sum = []
    theory_pair = []
    for ch in channels.tolist():
        if ch not in theory_sum_map or ch not in theory_pair_map:
            raise KeyError(f"Theory CSV missing channel l={ch}")
        theory_sum.append(theory_sum_map[ch])
        theory_pair.append(theory_pair_map[ch])

    return np.asarray(theory_sum, dtype=float), np.asarray(theory_pair, dtype=float)


def _set_common_axes(ax: plt.Axes, channels: np.ndarray) -> None:
    x = np.arange(channels.size)
    ax.set_xticks(x)
    ax.set_xticklabels([f"${int(c)}$" for c in channels])
    ax.set_xlabel(r"Target mode $l$")
    ax.set_ylabel(r"Extinction ratio (dB)")

    ax.spines["top"].set_visible(True)
    ax.spines["right"].set_visible(True)
    ax.tick_params(axis="both", which="major", direction="in", top=True, right=True)
    ax.yaxis.set_minor_locator(ticker.AutoMinorLocator(2))
    ax.tick_params(axis="y", which="minor", direction="in", right=True, length=2.0, width=0.6)

    ax.yaxis.grid(True, linestyle=":", linewidth=0.5, color="#cccccc", zorder=0)
    ax.set_axisbelow(True)


def plot_theory_vs_experiment(
    channels: np.ndarray,
    theory_vals: np.ndarray,
    exp_stats: dict[str, np.ndarray],
    mode: str,
    out_pdf: Path,
    title: str,
    theory_label: str,
    exp_label: str,
    dpi: int,
) -> None:
    exp_mean = exp_stats["mean"]
    x = np.arange(channels.size)
    width = 0.36

    fig, ax = plt.subplots(figsize=(5.2, 3.6), dpi=dpi)

    bars_theory = ax.bar(
        x - width / 2,
        theory_vals,
        width=width,
        color="#92C5DE",
        edgecolor="black",
        linewidth=0.6,
        alpha=0.92,
        label=theory_label,
        zorder=3,
    )

    bars_exp = ax.bar(
        x + width / 2,
        exp_mean,
        width=width,
        color="#2166AC",
        edgecolor="black",
        linewidth=0.6,
        alpha=0.92,
        label=exp_label,
        zorder=3,
    )

    if "scatter" in mode:
        raw = exp_stats["raw"]
        # n=3 points, add specific horizontal jitter
        jitter = np.array([-0.06, 0.0, 0.06])
        for i in range(channels.size):
            x_center = x[i] + width / 2
            ax.scatter(
                x_center + jitter,
                raw[i, :],
                s=20,
                facecolors="white",
                edgecolors="black",
                linewidths=0.7,
                alpha=0.85,
                zorder=5,
            )
            
    if mode != "scatter":
        err_key = "sd" if mode == "scatter_sd" else mode
        exp_err = get_yerr_for_mode(exp_stats, err_key)
        ax.errorbar(
            x + width / 2,
            exp_mean,
            yerr=exp_err,
            fmt="none",
            ecolor="black",
            elinewidth=0.9,
            capsize=2.8,
            capthick=0.9,
            zorder=4,
        )

    _set_common_axes(ax, channels)

    if mode in ("minmax", "scatter"):
        exp_lo = exp_stats["min"]
        exp_hi = exp_stats["max"]
    else:
        err_key = "sd" if mode == "scatter_sd" else mode
        exp_lo = exp_mean - exp_stats[err_key]
        exp_hi = exp_mean + exp_stats[err_key]

    y_min = min(float(np.min(theory_vals)), float(np.min(exp_lo)))
    y_max = max(float(np.max(theory_vals)), float(np.max(exp_hi)))
    lo = 0.0
    hi = math.ceil((y_max + 0.8) * 2) / 2
    ax.set_ylim(lo, hi)
    ax.set_xlim(-0.6, channels.size - 0.4)

    ax.set_title(title)
    ax.legend(
        handles=[bars_theory, bars_exp],
        loc="upper right",
        frameon=False,
        fontsize=8,
        handlelength=1.4,
        labelspacing=0.35,
    )

    fig.tight_layout(pad=0.45)
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_pdf, format="pdf", dpi=dpi, bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)


def resolve_theory_csv(path_from_arg: Path, script_dir: Path) -> Path:
    # Planned default path (as requested), then fallback to existing dataset location.
    if path_from_arg.exists():
        return path_from_arg

    fallback = (script_dir.parent.parent.parent / "理论tau和信噪比对比" / "tao_snr_l0_to_l8.csv").resolve()
    if fallback.exists():
        return fallback

    raise FileNotFoundError(
        f"Theory CSV not found. Tried:\n  1) {path_from_arg}\n  2) {fallback}"
    )


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description="Plot ER_sum and ER_pairmax theory-vs-experiment charts.")
    parser.add_argument(
        "--agg-xlsx",
        type=Path,
        default=script_dir.parent / "triplicate_full9_aggregate.xlsx",
        help="Path to triplicate aggregate workbook.",
    )
    parser.add_argument(
        "--theory-csv",
        type=Path,
        default=script_dir.parent.parent / "理论tau和信噪比对比" / "tao_snr_l0_to_l8.csv",
        help="Path to theory CSV.",
    )
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=script_dir / "输出图",
        help="Output directory.",
    )
    parser.add_argument(
        "--errorbar-mode",
        type=str,
        default="all",
        choices=["all", *ERRORBAR_MODES],
        help="Errorbar mode: all | ci95_t | sd | sem | minmax (default: all).",
    )
    parser.add_argument("--dpi", type=int, default=300, help="Figure DPI (default: 300).")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    script_dir = Path(__file__).resolve().parent

    if not args.agg_xlsx.exists():
        raise FileNotFoundError(f"Missing aggregate workbook: {args.agg_xlsx}")
    theory_csv = resolve_theory_csv(args.theory_csv, script_dir)

    channels, er_sum_runs, er_max_runs = parse_er_summary(args.agg_xlsx)
    er_sum_stats = compute_error_stats(er_sum_runs)
    er_max_stats = compute_error_stats(er_max_runs)

    theory_sum, theory_pair = parse_theory_csv(theory_csv, channels)

    configure_plot_style()
    modes = ERRORBAR_MODES if args.errorbar_mode == "all" else (args.errorbar_mode,)
    saved_paths: list[Path] = []
    for mode in modes:
        out1 = args.out_dir / f"ER_sum理论实验对比_含误差条_{MODE_TO_SUFFIX[mode]}.pdf"
        out2 = args.out_dir / f"ER_pairmax理论实验对比_含误差条_{MODE_TO_SUFFIX[mode]}.pdf"

        plot_theory_vs_experiment(
            channels=channels,
            theory_vals=theory_sum,
            exp_stats=er_sum_stats,
            mode=mode,
            out_pdf=out1,
            title=r"$ER_{\mathrm{sum}}$: Theory vs Experiment",
            theory_label=r"Theory",
            exp_label=r"Experiment",
            dpi=args.dpi,
        )

        plot_theory_vs_experiment(
            channels=channels,
            theory_vals=theory_pair,
            exp_stats=er_max_stats,
            mode=mode,
            out_pdf=out2,
            title=r"$ER_{\mathrm{pairmax}}$: Theory vs Experiment",
            theory_label=r"Theory",
            exp_label=r"Experiment",
            dpi=args.dpi,
        )
        saved_paths.extend([out1, out2])

        # Keep legacy output names mapped to CI95(t) for compatibility.
        if mode == "ci95_t":
            legacy1 = args.out_dir / "ER_sum理论实验对比_含误差条.pdf"
            legacy2 = args.out_dir / "ER_pairmax理论实验对比_含误差条.pdf"
            plot_theory_vs_experiment(
                channels=channels,
                theory_vals=theory_sum,
                exp_stats=er_sum_stats,
                mode=mode,
                out_pdf=legacy1,
                title=r"$ER_{\mathrm{sum}}$: Theory vs Experiment",
                theory_label=r"Theory",
                exp_label=r"Experiment",
                dpi=args.dpi,
            )
            plot_theory_vs_experiment(
                channels=channels,
                theory_vals=theory_pair,
                exp_stats=er_max_stats,
                mode=mode,
                out_pdf=legacy2,
                title=r"$ER_{\mathrm{pairmax}}$: Theory vs Experiment",
                theory_label=r"Theory",
                exp_label=r"Experiment",
                dpi=args.dpi,
            )
            saved_paths.extend([legacy1, legacy2])

    print(f"Theory CSV used: {theory_csv.resolve()}")
    for p in saved_paths:
        print(f"Saved: {p.resolve()}")


if __name__ == "__main__":
    main()
