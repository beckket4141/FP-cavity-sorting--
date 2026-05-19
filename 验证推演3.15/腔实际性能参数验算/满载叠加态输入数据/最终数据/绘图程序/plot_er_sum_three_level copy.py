#!/usr/bin/env python3
"""Plot the Chapter 5 ER_sum three-level comparison figure.

Data layers:
1. Experiment mean +- SD from triplicate_full9_aggregate.xlsx / ER_summary
2. Finite-M Airy theory on the ideal design branch k*=2/9 with measured F=32.21
3. Infinite-dimensional Lorentz conservative benchmark at tau0 = 3

Outputs:
  - 输出图/ER_sum_三级对比_含误差条.pdf
  - 输出图/ER_sum_三级对比_含误差条.png

Optional copy targets:
  - thesis/figures/ch05/ER_sum_三级对比_含误差条.pdf
  - thesis/figures/ch05/ER_sum_三级对比_含误差条.png
"""

from __future__ import annotations

import argparse
import math
import shutil
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import numpy as np
from openpyxl import load_workbook


F_EXP = 32.21
K_IDEAL = 2.0 / 9.0
ER_INF_LORENTZ = 10.469296343529393


def configure_plot_style() -> None:
    mpl.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["STIX Two Text", "STIXGeneral", "Times New Roman"],
            "mathtext.fontset": "stix",
            "font.size": 9,
            "axes.labelsize": 10,
            "axes.linewidth": 0.6,
            "axes.grid": False,
            "xtick.labelsize": 8.5,
            "ytick.labelsize": 8.5,
            "xtick.direction": "in",
            "ytick.direction": "in",
            "xtick.major.width": 0.5,
            "ytick.major.width": 0.5,
            "xtick.major.size": 3.0,
            "ytick.major.size": 3.0,
            "legend.fontsize": 7.8,
            "legend.frameon": False,
            "figure.dpi": 150,
        }
    )


def parse_er_summary(agg_xlsx: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    wb = load_workbook(agg_xlsx, read_only=True, data_only=True)
    if "ER_summary" not in wb.sheetnames:
        wb.close()
        raise KeyError(f"Missing sheet 'ER_summary' in {agg_xlsx}")

    ws = wb["ER_summary"]
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    idx = {str(v): i + 1 for i, v in enumerate(header) if v is not None}

    required = [
        "channel(lock=l)",
        "ER_sum_run1_dB",
        "ER_sum_run2_dB",
        "ER_sum_run3_dB",
    ]
    for key in required:
        if key not in idx:
            wb.close()
            raise KeyError(f"Missing column '{key}' in ER_summary")

    channels: list[int] = []
    runs: list[list[float]] = []
    for r in range(2, ws.max_row + 1):
        channel = ws.cell(r, idx["channel(lock=l)"]).value
        if channel is None:
            break
        channels.append(int(float(channel)))
        runs.append(
            [
                float(ws.cell(r, idx["ER_sum_run1_dB"]).value),
                float(ws.cell(r, idx["ER_sum_run2_dB"]).value),
                float(ws.cell(r, idx["ER_sum_run3_dB"]).value),
            ]
        )

    wb.close()
    order = np.argsort(np.asarray(channels, dtype=int))
    channels_arr = np.asarray(channels, dtype=int)[order]
    run_arr = np.asarray(runs, dtype=float)[order, :]
    means = np.mean(run_arr, axis=1)
    sds = np.std(run_arr, axis=1, ddof=1)
    return channels_arr, means, sds


def ideal_positions_from_k(channels: np.ndarray, k: float) -> np.ndarray:
    base = int(channels[0])
    return np.asarray([((int(ch) - base) * k) % 1.0 for ch in channels], dtype=float)


def geodesic_distance(a: float, b: float) -> float:
    d = abs(a - b)
    return min(d, 1.0 - d)


def airy_transmission(s: float, finesse: float) -> float:
    prefactor = (2.0 * finesse / math.pi) ** 2
    return 1.0 / (1.0 + prefactor * (math.sin(math.pi * s) ** 2))


def compute_airy_er_sum(channels: np.ndarray, positions: np.ndarray, finesse: float) -> np.ndarray:
    airy_er: list[float] = []
    for i, _ in enumerate(channels):
        total_leak = 0.0
        for j, _ in enumerate(channels):
            if i == j:
                continue
            s_ij = geodesic_distance(float(positions[i]), float(positions[j]))
            total_leak += airy_transmission(s_ij, finesse)
        airy_er.append(-10.0 * math.log10(total_leak))
    return np.asarray(airy_er, dtype=float)


def plot_three_level(
    channels: np.ndarray,
    exp_mean: np.ndarray,
    exp_sd: np.ndarray,
    airy_theory: np.ndarray,
    out_pdf: Path,
    out_png: Path,
    dpi: int,
) -> None:
    x = np.arange(channels.size, dtype=float)
    fig, ax = plt.subplots(figsize=(5.35, 3.95), dpi=dpi)

    airy_min = float(np.min(airy_theory))
    airy_max = float(np.max(airy_theory))

    exp = ax.errorbar(
        x,
        exp_mean,
        yerr=exp_sd,
        fmt="o",
        color="#2166AC",
        ecolor="#2166AC",
        markerfacecolor="#2166AC",
        markeredgecolor="black",
        markeredgewidth=0.5,
        markersize=4.8,
        elinewidth=1.2,
        capsize=2.8,
        capthick=1.2,
        zorder=4,
        label="Experiment",
    )

    airy_line = ax.axhline(
        np.mean(airy_theory),
        linestyle="-",
        linewidth=1.5,
        color="#ca0020",
        alpha=0.8,
        zorder=2.6,
        label=r"Finite-$M$ Airy theory ($k^*=2/9$)",
    )

    lower = ax.axhline(
        ER_INF_LORENTZ,
        color="#666666",
        linestyle=(0, (4, 2)),
        linewidth=1.0,
        zorder=2,
        label=r"Lorentz conservative benchmark ($10.47\ \mathrm{dB}$)",
    )

    ax.set_xticks(x)
    ax.set_xticklabels([rf"${int(c)}$" for c in channels])
    ax.set_xlabel(r"Target mode $l$")
    ax.set_ylabel(r"$ER_{\mathrm{sum}}$ (dB)")

    ax.tick_params(axis="both", which="major", direction="in", top=True, right=True)
    ax.yaxis.set_minor_locator(ticker.AutoMinorLocator(2))
    ax.tick_params(axis="y", which="minor", direction="in", right=True, length=2.0, width=0.5)
    ax.yaxis.grid(True, linestyle=":", linewidth=0.5, color="#cccccc", zorder=0)
    ax.set_axisbelow(True)

    y_min = min(float(np.min(exp_mean - exp_sd)), airy_min, ER_INF_LORENTZ)
    y_max = max(float(np.max(exp_mean + exp_sd)), airy_max, ER_INF_LORENTZ)
    ax.set_ylim(math.floor((y_min - 0.3) * 2) / 2, math.ceil((y_max + 0.35) * 2) / 2)
    ax.set_xlim(-0.45, channels.size - 0.55)

    ax.legend(
        handles=[exp, airy_line, lower],
        loc="best",
        frameon=False,
        fontsize=7.8,
        handlelength=1.8,
        labelspacing=0.4,
    )

    fig.tight_layout(pad=0.45)
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_pdf, format="pdf", dpi=dpi, bbox_inches="tight", pad_inches=0.03)
    fig.savefig(out_png, format="png", dpi=dpi, bbox_inches="tight", pad_inches=0.03)
    plt.close(fig)


def parse_args() -> argparse.Namespace:
    script_dir = Path(__file__).resolve().parent
    base_dir = script_dir.parent
    thesis_fig_dir = script_dir.parent.parent.parent.parent.parent / "thesis" / "figures" / "ch05"

    parser = argparse.ArgumentParser(description="Plot ER_sum three-level comparison for Chapter 5.")
    parser.add_argument(
        "--agg-xlsx",
        type=Path,
        default=base_dir / "triplicate_full9_aggregate.xlsx",
        help="Triplicate aggregate workbook.",
    )
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=script_dir / "输出图",
        help="Directory for generated figures.",
    )
    parser.add_argument(
        "--copy-to-thesis",
        action="store_true",
        help="Also copy the generated PDF/PNG into thesis/figures/ch05.",
    )
    parser.add_argument(
        "--thesis-fig-dir",
        type=Path,
        default=thesis_fig_dir,
        help="Target thesis figure directory when --copy-to-thesis is enabled.",
    )
    parser.add_argument("--dpi", type=int, default=300, help="Figure DPI.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    configure_plot_style()

    channels_exp, exp_mean, exp_sd = parse_er_summary(args.agg_xlsx)
    positions = ideal_positions_from_k(channels_exp, K_IDEAL)
    airy_theory = compute_airy_er_sum(channels_exp, positions, F_EXP)

    out_pdf = args.out_dir / "ER_sum_三级对比_含误差条.pdf"
    out_png = args.out_dir / "ER_sum_三级对比_含误差条.png"
    plot_three_level(channels_exp, exp_mean, exp_sd, airy_theory, out_pdf, out_png, args.dpi)

    print(f"Saved PDF: {out_pdf.resolve()}")
    print(f"Saved PNG: {out_png.resolve()}")
    print(f"Airy theory branch: k* = {K_IDEAL:.12f} (ideal design branch)")
    print(f"Airy theory finesse: F = {F_EXP:.2f}")
    print(f"Finite-M Airy mean ER_sum: {np.mean(airy_theory):.6f} dB")
    print(f"Finite-M Airy min ER_sum: {np.min(airy_theory):.6f} dB")
    print(f"Finite-M Airy spread: {(np.max(airy_theory) - np.min(airy_theory)):.6f} dB")

    if args.copy_to_thesis:
        args.thesis_fig_dir.mkdir(parents=True, exist_ok=True)
        thesis_pdf = args.thesis_fig_dir / out_pdf.name
        thesis_png = args.thesis_fig_dir / out_png.name
        shutil.copy2(out_pdf, thesis_pdf)
        shutil.copy2(out_png, thesis_png)
        print(f"Copied PDF: {thesis_pdf.resolve()}")
        print(f"Copied PNG: {thesis_png.resolve()}")


if __name__ == "__main__":
    main()
