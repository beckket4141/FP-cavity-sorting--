#!/usr/bin/env python3
"""Compare full-matrix de-embedding against diagonal-only correction.

This script keeps the original data files untouched. It re-computes the 9D
full-load result for each run from the raw source files:

  - final_matrix_summary.xlsx
  - transmission workbook (usually named 透射率.xlsx)

Two correction modes are compared on exactly the same raw observations:

1. Full matrix:
     y = S x, solved by NNLS with x >= 0
2. Diagonal only:
     y = D x, where D = diag(S)
     This is equivalent to per-channel efficiency correction only.

The report focuses on:
  - percentage matrix change
  - ER_sum / ER_max change
  - signal-share change
  - threshold crossing relative to the thesis 10.47 dB design floor
"""

from __future__ import annotations

import argparse
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import nnls


LOCK_RE = re.compile(r"lock\s*=\s*(\d+)", re.IGNORECASE)
CH_RE = re.compile(r"ch\s*=\s*(\d+)", re.IGNORECASE)
L_RE = re.compile(r"l\s*=\s*(\d+)", re.IGNORECASE)


@dataclass
class MatrixData:
    channels: list[int]
    locks: list[int]
    values: np.ndarray


@dataclass
class MatrixBundle:
    mean: MatrixData
    std: MatrixData
    var: MatrixData
    n: MatrixData


@dataclass
class ChannelMetrics:
    er_sum_db: dict[int, float]
    er_max_db: dict[int, float]
    signal_share: dict[int, float]


@dataclass
class RunComparison:
    run_name: str
    channels: list[int]
    locks: list[int]
    s_matrix: np.ndarray
    y_matrix: np.ndarray
    x_full: np.ndarray
    x_diag: np.ndarray
    pct_full: np.ndarray
    pct_diag: np.ndarray
    metrics_full: ChannelMetrics
    metrics_diag: ChannelMetrics
    cond_s: float
    diag_eff_min: float
    diag_eff_max: float
    mean_offdiag_to_diag: float
    max_offdiag_to_diag: float


def parse_float(v) -> float:
    if v is None:
        return float("nan")
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    s = s.replace("_x000d_", "").strip()
    if s in {"", "NaN", "nan"}:
        return float("nan")
    if s == "Inf":
        return float("inf")
    if s == "-Inf":
        return float("-inf")
    try:
        return float(s)
    except ValueError:
        m = re.search(r"[-+]?\d*\.?\d+(?:[eE][-+]?\d+)?", s)
        if m:
            return float(m.group(0))
        raise


def safe_db10(num: float, den: float) -> float:
    if not math.isfinite(num) or not math.isfinite(den) or num <= 0 or den <= 0:
        return float("nan")
    return 10.0 * math.log10(num / den)


def parse_matrix_sheet(wb, sheet_name: str) -> MatrixData:
    if sheet_name not in wb.sheetnames:
        raise KeyError(f"Missing sheet: {sheet_name}")
    ws = wb[sheet_name]

    locks: list[int] = []
    for cell in ws[1][1:]:
        if cell.value is None:
            continue
        m = LOCK_RE.search(str(cell.value))
        if m:
            locks.append(int(m.group(1)))

    channels: list[int] = []
    rows: list[list[float]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        m = CH_RE.search(str(row[0]))
        if not m:
            continue
        channels.append(int(m.group(1)))
        rows.append([parse_float(v) for v in row[1 : 1 + len(locks)]])

    if not channels or not locks:
        raise ValueError(f"Sheet {sheet_name} has no valid matrix data")

    return MatrixData(channels=channels, locks=locks, values=np.asarray(rows, dtype=float))


def parse_bundle(wb, prefix: str) -> MatrixBundle:
    mean = parse_matrix_sheet(wb, f"{prefix}_mean")
    std = parse_matrix_sheet(wb, f"{prefix}_std")
    var = parse_matrix_sheet(wb, f"{prefix}_var")
    n = parse_matrix_sheet(wb, f"{prefix}_n")

    for other_name, other in [("std", std), ("var", var), ("n", n)]:
        if mean.channels != other.channels or mean.locks != other.locks:
            raise ValueError(f"{prefix}_{other_name} axis mismatch")

    return MatrixBundle(mean=mean, std=std, var=var, n=n)


def parse_transmission_table(trans_xlsx: Path) -> dict[int, float]:
    wb = load_workbook(trans_xlsx, data_only=True, read_only=True)
    ws = None
    for candidate in wb.worksheets:
        if candidate.max_row >= 2:
            ws = candidate
            break
    if ws is None:
        wb.close()
        raise ValueError(f"No usable sheet in {trans_xlsx.name}")

    out: dict[int, float] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        label = row[0]
        if label is None:
            continue
        m = L_RE.search(str(label))
        if not m:
            continue
        lock = int(m.group(1))
        p_after_uW = parse_float(row[3]) if len(row) > 3 else float("nan")
        if math.isfinite(p_after_uW) and p_after_uW > 0:
            out[lock] = p_after_uW * 1e-6

    wb.close()

    if not out:
        raise ValueError(f"No valid transmission rows parsed from {trans_xlsx.name}")
    return out


def normalize_calibration_by_trans(calib: MatrixBundle, trans_watts: dict[int, float]) -> MatrixBundle:
    mean = calib.mean.values.copy()
    std = calib.std.values.copy()
    var = calib.var.values.copy()
    n = calib.n.values.copy()

    for col, lock in enumerate(calib.mean.locks):
        if lock not in trans_watts:
            raise KeyError(f"Missing transmission entry for lock={lock}")
        scale = trans_watts[lock]
        mean[:, col] = mean[:, col] / scale
        std[:, col] = std[:, col] / scale
        var[:, col] = var[:, col] / (scale ** 2)

    return MatrixBundle(
        mean=MatrixData(calib.mean.channels, calib.mean.locks, mean),
        std=MatrixData(calib.std.channels, calib.std.locks, std),
        var=MatrixData(calib.var.channels, calib.var.locks, var),
        n=MatrixData(calib.n.channels, calib.n.locks, n),
    )


def subset_bundle(bundle: MatrixBundle, channels: list[int], locks: list[int]) -> MatrixBundle:
    ch_to_row = {ch: idx for idx, ch in enumerate(bundle.mean.channels)}
    lk_to_col = {lk: idx for idx, lk in enumerate(bundle.mean.locks)}
    row_idx = [ch_to_row[ch] for ch in channels]
    col_idx = [lk_to_col[lk] for lk in locks]

    def subset(md: MatrixData) -> MatrixData:
        return MatrixData(
            channels=channels[:],
            locks=locks[:],
            values=md.values[np.ix_(row_idx, col_idx)],
        )

    return MatrixBundle(
        mean=subset(bundle.mean),
        std=subset(bundle.std),
        var=subset(bundle.var),
        n=subset(bundle.n),
    )


def find_transmission_xlsx(run_dir: Path) -> Path:
    preferred = run_dir / "透射率.xlsx"
    if preferred.exists():
        return preferred

    known = {
        "final_matrix_summary.xlsx",
        "calib_matrix_with_uncertainty.xlsx",
        "full9_matrix_with_uncertainty.xlsx",
        "full4_matrix_with_uncertainty.xlsx",
        "channel_snr_analysis.xlsx",
    }
    candidates = [p for p in run_dir.glob("*.xlsx") if p.name not in known]
    if len(candidates) == 1:
        return candidates[0]
    raise FileNotFoundError(f"Could not uniquely determine transmission workbook in {run_dir}")


def solve_full_nnls(s_matrix: np.ndarray, y_matrix: np.ndarray) -> np.ndarray:
    x = np.zeros_like(y_matrix, dtype=float)
    for col in range(y_matrix.shape[1]):
        y = np.clip(np.nan_to_num(y_matrix[:, col], nan=0.0), 0.0, None)
        sol, _ = nnls(s_matrix, y)
        x[:, col] = sol
    return x


def solve_diag_only(s_matrix: np.ndarray, y_matrix: np.ndarray, channels: list[int], locks: list[int]) -> np.ndarray:
    lock_to_col = {lock: idx for idx, lock in enumerate(locks)}
    diag_eff = []
    for row_idx, ch in enumerate(channels):
        if ch not in lock_to_col:
            raise ValueError(f"Channel {ch} has no matching lock; diagonal-only model is undefined")
        val = float(s_matrix[row_idx, lock_to_col[ch]])
        if not math.isfinite(val) or val <= 0:
            raise ValueError(f"Invalid diagonal efficiency for channel {ch}: {val}")
        diag_eff.append(val)
    diag_eff_arr = np.asarray(diag_eff, dtype=float)[:, None]
    return np.divide(np.clip(np.nan_to_num(y_matrix, nan=0.0), 0.0, None), diag_eff_arr)


def normalize_columns(x_matrix: np.ndarray) -> np.ndarray:
    col_sums = np.sum(x_matrix, axis=0, keepdims=True)
    return np.divide(x_matrix, col_sums, out=np.zeros_like(x_matrix), where=col_sums > 0)


def compute_metrics(x_matrix: np.ndarray, channels: list[int], locks: list[int]) -> ChannelMetrics:
    ch_to_row = {ch: idx for idx, ch in enumerate(channels)}
    er_sum_db: dict[int, float] = {}
    er_max_db: dict[int, float] = {}
    signal_share: dict[int, float] = {}

    for col, lock in enumerate(locks):
        if lock not in ch_to_row:
            continue
        sig_row = ch_to_row[lock]
        col_vec = np.clip(x_matrix[:, col], 0.0, None)
        signal = float(col_vec[sig_row])
        leak = np.delete(col_vec, sig_row)
        leak_sum = float(np.sum(leak)) if leak.size else float("nan")
        leak_max = float(np.max(leak)) if leak.size else float("nan")
        total = float(np.sum(col_vec))

        er_sum_db[lock] = safe_db10(signal, leak_sum)
        er_max_db[lock] = safe_db10(signal, leak_max)
        signal_share[lock] = signal / total if total > 0 else float("nan")

    return ChannelMetrics(er_sum_db=er_sum_db, er_max_db=er_max_db, signal_share=signal_share)


def load_run(run_dir: Path) -> RunComparison:
    matrix_xlsx = run_dir / "final_matrix_summary.xlsx"
    trans_xlsx = find_transmission_xlsx(run_dir)
    if not matrix_xlsx.exists():
        raise FileNotFoundError(f"Missing {matrix_xlsx}")

    wb = load_workbook(matrix_xlsx, data_only=True, read_only=True)
    calib_raw = parse_bundle(wb, "calib")
    full9_raw = parse_bundle(wb, "full9")
    wb.close()

    trans_watts = parse_transmission_table(trans_xlsx)
    calib_norm = normalize_calibration_by_trans(calib_raw, trans_watts)
    calib_9 = subset_bundle(calib_norm, full9_raw.mean.channels, full9_raw.mean.locks)

    channels = full9_raw.mean.channels
    locks = full9_raw.mean.locks
    s_matrix = calib_9.mean.values
    y_matrix = full9_raw.mean.values

    x_full = solve_full_nnls(s_matrix, y_matrix)
    x_diag = solve_diag_only(s_matrix, y_matrix, channels, locks)

    pct_full = normalize_columns(x_full)
    pct_diag = normalize_columns(x_diag)

    metrics_full = compute_metrics(x_full, channels, locks)
    metrics_diag = compute_metrics(x_diag, channels, locks)

    lock_to_col = {lock: idx for idx, lock in enumerate(locks)}
    diag_eff = np.asarray([s_matrix[row, lock_to_col[ch]] for row, ch in enumerate(channels)], dtype=float)

    offdiag_ratios: list[float] = []
    for col, lock in enumerate(locks):
        sig_row = channels.index(lock)
        diag_val = float(s_matrix[sig_row, col])
        offdiag_sum = float(np.sum(s_matrix[:, col]) - diag_val)
        offdiag_ratios.append(offdiag_sum / diag_val if diag_val > 0 else float("nan"))

    return RunComparison(
        run_name=run_dir.name,
        channels=channels,
        locks=locks,
        s_matrix=s_matrix,
        y_matrix=y_matrix,
        x_full=x_full,
        x_diag=x_diag,
        pct_full=pct_full,
        pct_diag=pct_diag,
        metrics_full=metrics_full,
        metrics_diag=metrics_diag,
        cond_s=float(np.linalg.cond(s_matrix)),
        diag_eff_min=float(np.min(diag_eff)),
        diag_eff_max=float(np.max(diag_eff)),
        mean_offdiag_to_diag=float(np.nanmean(offdiag_ratios)),
        max_offdiag_to_diag=float(np.nanmax(offdiag_ratios)),
    )


def finite_stats(values: Iterable[float]) -> tuple[float, float, float]:
    arr = np.asarray(list(values), dtype=float)
    arr = arr[np.isfinite(arr)]
    if arr.size == 0:
        return float("nan"), float("nan"), float("nan")
    return float(np.mean(arr)), float(np.min(arr)), float(np.max(arr))


def mean_std(values: Iterable[float]) -> tuple[float, float]:
    arr = np.asarray(list(values), dtype=float)
    arr = arr[np.isfinite(arr)]
    if arr.size == 0:
        return float("nan"), float("nan")
    if arr.size == 1:
        return float(arr[0]), 0.0
    return float(np.mean(arr)), float(np.std(arr, ddof=1))


def format_matrix_percent(mat: np.ndarray, channels: list[int], locks: list[int]) -> str:
    headers = ["ch\\lock"] + [str(lock) for lock in locks]
    rows = []
    for row_idx, ch in enumerate(channels):
        row = [str(ch)] + [f"{100.0 * mat[row_idx, col]:6.2f}" for col in range(len(locks))]
        rows.append(row)

    widths = [max(len(headers[col]), max(len(row[col]) for row in rows)) for col in range(len(headers))]
    out_lines = []
    out_lines.append("  ".join(headers[col].rjust(widths[col]) for col in range(len(headers))))
    for row in rows:
        out_lines.append("  ".join(row[col].rjust(widths[col]) for col in range(len(row))))
    return "\n".join(out_lines)


def top_element_changes(
    pct_full: np.ndarray,
    pct_diag: np.ndarray,
    channels: list[int],
    locks: list[int],
    top_k: int,
) -> list[str]:
    diff = 100.0 * (pct_diag - pct_full)
    flat = []
    for row_idx, ch in enumerate(channels):
        for col_idx, lock in enumerate(locks):
            flat.append((abs(diff[row_idx, col_idx]), ch, lock, diff[row_idx, col_idx], 100.0 * pct_full[row_idx, col_idx], 100.0 * pct_diag[row_idx, col_idx]))
    flat.sort(reverse=True)
    lines = []
    for _, ch, lock, delta, full_val, diag_val in flat[:top_k]:
        lines.append(
            f"  ch={ch}, lock={lock}: full={full_val:.2f}%, diag-only={diag_val:.2f}%, delta={delta:+.2f} pct-pt"
        )
    return lines


def build_report(run_results: list[RunComparison], threshold_db: float, top_k: int) -> str:
    if not run_results:
        return "No runs loaded."

    channels = run_results[0].channels
    locks = run_results[0].locks

    lines: list[str] = []
    lines.append("Diagonal-only vs full-matrix comparison for 9D full-load data")
    lines.append(f"Runs: {', '.join(result.run_name for result in run_results)}")
    lines.append(f"Design floor checked: ER_sum >= {threshold_db:.2f} dB")
    lines.append("")

    for result in run_results:
        er_full_vals = [result.metrics_full.er_sum_db[lock] for lock in locks]
        er_diag_vals = [result.metrics_diag.er_sum_db[lock] for lock in locks]
        er_drop = [result.metrics_diag.er_sum_db[lock] - result.metrics_full.er_sum_db[lock] for lock in locks]
        share_full_vals = [100.0 * result.metrics_full.signal_share[lock] for lock in locks]
        share_diag_vals = [100.0 * result.metrics_diag.signal_share[lock] for lock in locks]
        share_drop = [d - f for d, f in zip(share_diag_vals, share_full_vals)]

        er_full_mean, er_full_min, er_full_max = finite_stats(er_full_vals)
        er_diag_mean, er_diag_min, er_diag_max = finite_stats(er_diag_vals)
        drop_mean, drop_min, drop_max = finite_stats(er_drop)
        share_full_mean, _, _ = finite_stats(share_full_vals)
        share_diag_mean, _, _ = finite_stats(share_diag_vals)
        share_drop_mean, share_drop_min, share_drop_max = finite_stats(share_drop)

        pct_diff = 100.0 * (result.pct_diag - result.pct_full)
        mean_abs_pct_diff = float(np.mean(np.abs(pct_diff)))
        max_abs_pct_diff = float(np.max(np.abs(pct_diff)))

        below_full = sum(v < threshold_db for v in er_full_vals if math.isfinite(v))
        below_diag = sum(v < threshold_db for v in er_diag_vals if math.isfinite(v))

        lines.append(f"[Run {result.run_name}]")
        lines.append(
            f"  cond(S)={result.cond_s:.3f}; diag efficiency range="
            f"[{100.0 * result.diag_eff_min:.2f}%, {100.0 * result.diag_eff_max:.2f}%]; "
            f"mean(offdiag/diag)={100.0 * result.mean_offdiag_to_diag:.2f}%, "
            f"max(offdiag/diag)={100.0 * result.max_offdiag_to_diag:.2f}%"
        )
        lines.append(
            f"  ER_sum full={er_full_mean:.3f} dB (min {er_full_min:.3f}, max {er_full_max:.3f}); "
            f"diag-only={er_diag_mean:.3f} dB (min {er_diag_min:.3f}, max {er_diag_max:.3f}); "
            f"delta={drop_mean:+.3f} dB on average"
        )
        lines.append(
            f"  Signal share full={share_full_mean:.2f}%; diag-only={share_diag_mean:.2f}%; "
            f"delta={share_drop_mean:+.2f} pct-pt on average"
        )
        lines.append(
            f"  Percentage-matrix change: mean |delta|={mean_abs_pct_diff:.2f} pct-pt, "
            f"max |delta|={max_abs_pct_diff:.2f} pct-pt"
        )
        lines.append(
            f"  Channels below {threshold_db:.2f} dB: full={below_full}/{len(locks)}, "
            f"diag-only={below_diag}/{len(locks)}"
        )

        channel_drop = sorted(
            (
                (
                    result.metrics_diag.er_sum_db[lock] - result.metrics_full.er_sum_db[lock],
                    lock,
                    result.metrics_full.er_sum_db[lock],
                    result.metrics_diag.er_sum_db[lock],
                    100.0 * result.metrics_full.signal_share[lock],
                    100.0 * result.metrics_diag.signal_share[lock],
                )
                for lock in locks
            ),
            key=lambda item: item[0],
        )
        lines.append("  Worst ER_sum drops:")
        for delta, lock, full_er, diag_er, full_share, diag_share in channel_drop[:top_k]:
            lines.append(
                f"    lock={lock}: ER_sum {full_er:.3f} -> {diag_er:.3f} dB "
                f"({delta:+.3f} dB), signal share {full_share:.2f}% -> {diag_share:.2f}%"
            )
        lines.append("")

    combined_pct_full = np.mean(np.stack([result.pct_full for result in run_results], axis=0), axis=0)
    combined_pct_diag = np.mean(np.stack([result.pct_diag for result in run_results], axis=0), axis=0)
    combined_x_full = np.mean(np.stack([result.x_full for result in run_results], axis=0), axis=0)
    combined_x_diag = np.mean(np.stack([result.x_diag for result in run_results], axis=0), axis=0)

    combined_metrics_full = compute_metrics(combined_x_full, channels, locks)
    combined_metrics_diag = compute_metrics(combined_x_diag, channels, locks)

    er_full_by_lock = {lock: [result.metrics_full.er_sum_db[lock] for result in run_results] for lock in locks}
    er_diag_by_lock = {lock: [result.metrics_diag.er_sum_db[lock] for result in run_results] for lock in locks}
    share_full_by_lock = {lock: [100.0 * result.metrics_full.signal_share[lock] for result in run_results] for lock in locks}
    share_diag_by_lock = {lock: [100.0 * result.metrics_diag.signal_share[lock] for result in run_results] for lock in locks}

    er_full_all = [result.metrics_full.er_sum_db[lock] for result in run_results for lock in locks]
    er_diag_all = [result.metrics_diag.er_sum_db[lock] for result in run_results for lock in locks]
    share_full_all = [100.0 * result.metrics_full.signal_share[lock] for result in run_results for lock in locks]
    share_diag_all = [100.0 * result.metrics_diag.signal_share[lock] for result in run_results for lock in locks]

    er_full_mean, er_full_std = mean_std(er_full_all)
    er_diag_mean, er_diag_std = mean_std(er_diag_all)
    share_full_mean, share_full_std = mean_std(share_full_all)
    share_diag_mean, share_diag_std = mean_std(share_diag_all)

    pct_diff_combined = 100.0 * (combined_pct_diag - combined_pct_full)
    mean_abs_pct_diff_combined = float(np.mean(np.abs(pct_diff_combined)))
    max_abs_pct_diff_combined = float(np.max(np.abs(pct_diff_combined)))

    below_full_all = sum(v < threshold_db for v in er_full_all if math.isfinite(v))
    below_diag_all = sum(v < threshold_db for v in er_diag_all if math.isfinite(v))

    lines.append("[Combined over 3 runs]")
    lines.append(
        f"  ER_sum full={er_full_mean:.3f} +/- {er_full_std:.3f} dB; "
        f"diag-only={er_diag_mean:.3f} +/- {er_diag_std:.3f} dB; "
        f"mean delta={er_diag_mean - er_full_mean:+.3f} dB"
    )
    lines.append(
        f"  Signal share full={share_full_mean:.2f} +/- {share_full_std:.2f}%; "
        f"diag-only={share_diag_mean:.2f} +/- {share_diag_std:.2f}%; "
        f"mean delta={share_diag_mean - share_full_mean:+.2f} pct-pt"
    )
    lines.append(
        f"  Percentage-matrix change: mean |delta|={mean_abs_pct_diff_combined:.2f} pct-pt, "
        f"max |delta|={max_abs_pct_diff_combined:.2f} pct-pt"
    )
    lines.append(
        f"  Threshold failures across all 27 channel-cases: full={below_full_all}/27, "
        f"diag-only={below_diag_all}/27"
    )
    lines.append("")
    lines.append("  Per-channel ER_sum summary (mean +/- std across runs):")
    for lock in locks:
        full_mean, full_std = mean_std(er_full_by_lock[lock])
        diag_mean, diag_std = mean_std(er_diag_by_lock[lock])
        share_full_mean_lock, share_full_std_lock = mean_std(share_full_by_lock[lock])
        share_diag_mean_lock, share_diag_std_lock = mean_std(share_diag_by_lock[lock])
        lines.append(
            f"    lock={lock}: ER_sum full={full_mean:.3f} +/- {full_std:.3f} dB, "
            f"diag-only={diag_mean:.3f} +/- {diag_std:.3f} dB, "
            f"delta={diag_mean - full_mean:+.3f} dB; "
            f"signal share full={share_full_mean_lock:.2f} +/- {share_full_std_lock:.2f}%, "
            f"diag-only={share_diag_mean_lock:.2f} +/- {share_diag_std_lock:.2f}%"
        )
    lines.append("")
    lines.append("  Combined percentage matrix from full-matrix correction (%):")
    lines.append(format_matrix_percent(combined_pct_full, channels, locks))
    lines.append("")
    lines.append("  Combined percentage matrix from diagonal-only correction (%):")
    lines.append(format_matrix_percent(combined_pct_diag, channels, locks))
    lines.append("")
    lines.append("  Largest combined matrix-element changes:")
    lines.extend(top_element_changes(combined_pct_full, combined_pct_diag, channels, locks, top_k))
    lines.append("")
    lines.append(
        "  Combined deterministic ER_sum from averaged matrices: "
        + ", ".join(
            f"lock={lock}: full={combined_metrics_full.er_sum_db[lock]:.3f} dB, "
            f"diag-only={combined_metrics_diag.er_sum_db[lock]:.3f} dB"
            for lock in locks
        )
    )

    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Compare diagonal-only correction against full-matrix NNLS correction."
    )
    parser.add_argument(
        "--base-dir",
        type=Path,
        default=Path(__file__).resolve().parent,
        help="Base directory that contains run folders (default: script directory).",
    )
    parser.add_argument(
        "--runs",
        nargs="+",
        default=["1", "2", "3"],
        help="Run folder names to analyze (default: 1 2 3).",
    )
    parser.add_argument(
        "--threshold-db",
        type=float,
        default=10.47,
        help="ER_sum design floor used in the report (default: 10.47 dB).",
    )
    parser.add_argument(
        "--top-k",
        type=int,
        default=3,
        help="How many worst channels / biggest matrix changes to print (default: 3).",
    )
    args = parser.parse_args()

    run_results = [load_run(args.base_dir / run_name) for run_name in args.runs]
    report = build_report(run_results, threshold_db=args.threshold_db, top_k=args.top_k)
    print(report)


if __name__ == "__main__":
    main()
