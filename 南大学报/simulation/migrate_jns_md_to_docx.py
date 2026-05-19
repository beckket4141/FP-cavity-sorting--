from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path.cwd()
MANUSCRIPT_DIR = next(ROOT.rglob("nju_jns_draft.md")).parent
MD_PATH = MANUSCRIPT_DIR / "nju_jns_draft.md"
DOCX_PATH = MANUSCRIPT_DIR / "nju_jns_template_working_copy.docx"
FIG_DIR = next(ROOT.rglob("fig01_unified_framework.png")).parent


TITLE_EN = (
    "Unified design framework and error mechanism for Laguerre-Gaussian "
    "mode sorting in stable two-mirror Fabry-Perot cavities"
)

ABSTRACT_EN = (
    "Fabry-Perot cavities can realize state-preserving spectral sorting of "
    "Laguerre-Gaussian modes through Gouy-phase-dependent resonance. Existing "
    "analyses often start from the plane-concave configuration or from a "
    "specific experimental implementation, while a unified design description "
    "for general stable two-mirror cavities is still needed. Based on a "
    "single-free-spectral-range folded-spectrum representation, this paper "
    "reduces mode-sorting design to a circular arrangement problem controlled "
    "by the effective normalized Gouy step. The cavity geometry enters through "
    "the mapping between the effective step and the geometric parameters. The "
    "results show that the plane-concave cavity obeys L/R=sin^2(pi k_eff), "
    "whereas the symmetric double-concave cavity obeys L/R=1+/-cos(pi k_eff); "
    "the latter cannot be obtained from the former by the simple substitution "
    "L to 2L. In the small-geometry-ratio limit, the step enhancement of the "
    "symmetric double-concave cavity relative to the plane-concave cavity tends "
    "to sqrt(2). The minimum-spacing landscape for consecutive mode sets shows "
    "a hierarchical refinement structure governed by neighboring Farey "
    "fractions, and different cavity geometries change the geometric "
    "back-substitution and robustness ranking of the optimal branches. For "
    "the M=9 consecutive mode set, the robust representative branch is m=2 in "
    "the plane-concave cavity, while it shifts to m=4 in the symmetric "
    "double-concave cavity. In the error mechanism, axisymmetric waist-size, "
    "waist-position, and curvature mismatches mainly excite higher radial "
    "orders within the same azimuthal-index subspace, with p=1 being the "
    "leading observable parasitic term. Symmetry-breaking errors such as "
    "lateral displacement and tilt are required to introduce appreciable "
    "azimuthal-mode mixing. These results provide a unified explanation for "
    "satellite peaks in measured spectra and indicate that LG mode sorting in "
    "stable two-mirror FP cavities should be optimized at both the effective "
    "step design level and the spatial-mode error-control level."
)

KEYWORDS_EN = (
    "Fabry-Perot cavity; Laguerre-Gaussian mode; Gouy phase; mode sorting; "
    "spectral folding; error mechanism"
)


FIGURES = {
    1: ("fig01_unified_framework.png", 14.0),
    2: ("fig02_landscape_structure.png", 14.0),
    3: ("fig03_geometry_robustness.png", 14.0),
    4: ("fig04_radial_leakage.png", 14.0),
    5: ("fig05_error_classification.png", 14.0),
}


def clear_document(doc: Document) -> None:
    body = doc._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_page(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.7)
    section.right_margin = Cm(2.7)
    section.start_type = WD_SECTION_START.CONTINUOUS
    for p in section.footer.paragraphs:
        p.clear()
    footer_lines = [
        "基金项目：待补",
        "收稿日期：此处请留空",
        "*通信联系人，E-mail：待补",
    ]
    for i, line in enumerate(footer_lines):
        p = section.footer.paragraphs[0] if i == 0 else section.footer.add_paragraph()
        add_run(p, line, size=8, east="宋体", latin="Times New Roman")


def add_run(
    p,
    text: str,
    *,
    size: float = 10.5,
    bold: bool = False,
    italic: bool = False,
    color: RGBColor | None = None,
    east: str = "宋体",
    latin: str = "Times New Roman",
):
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    font = run.font
    font.name = latin
    font.size = Pt(size)
    if color is not None:
        font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    return run


def style_paragraph(
    p,
    *,
    align: WD_ALIGN_PARAGRAPH | None = None,
    first_line: bool = False,
    before: float = 0,
    after: float = 0,
    line: float = 1.25,
) -> None:
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if first_line:
        fmt.first_line_indent = Pt(21)
    else:
        fmt.first_line_indent = None


def add_plain_paragraph(doc: Document, text: str, *, first_line: bool = True) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, first_line=first_line, after=3, line=1.25)
    add_inline_math_runs(p, text)


def add_inline_math_runs(p, text: str, *, size: float = 10.5) -> None:
    parts = re.split(r"(\\\(.+?\\\))", text)
    for part in parts:
        if not part:
            continue
        if part.startswith(r"\(") and part.endswith(r"\)"):
            add_run(p, latex_to_plain(part[2:-2]), size=size, latin="Cambria Math", east="Cambria Math")
        else:
            add_run(p, part, size=size)


def add_heading(doc: Document, text: str, level: int) -> None:
    style = "科技导报-一级标题" if level == 1 else "科技导报-2级标题"
    p = doc.add_paragraph(style=style if style in [s.name for s in doc.styles] else None)
    style_paragraph(p, before=8 if level == 1 else 5, after=4, line=1.2)
    add_run(p, text, size=12 if level == 1 else 10.5, bold=True, east="黑体")


def add_equation(doc: Document, equation_lines: list[str]) -> None:
    text = latex_to_plain(" ".join(line.strip() for line in equation_lines).strip())
    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=3, line=1.1)
    add_run(p, text, size=10.5, latin="Cambria Math", east="Cambria Math")


def latex_to_plain(text: str) -> str:
    """Convert the limited LaTeX used in the Markdown draft to readable Word text."""
    s = text.strip()
    s = re.sub(
        r"\\frac\{k_\{\\mathrm\{cc\}\}\}\{k_\{\\mathrm\{pc\}\}\}",
        "k_cc/k_pc",
        s,
    )
    s = s.replace(r"\,", " ")
    s = s.replace(r"\left", "").replace(r"\right", "")
    s = s.replace(r"\rightarrow", "→")
    s = s.replace(r"\pm", "±")
    s = s.replace(r"\approx", "≈")
    s = s.replace(r"\ge", "≥").replace(r"\le", "≤")
    s = s.replace(r"\neq", "≠")
    s = s.replace(r"\infty", "∞")
    s = s.replace(r"\{", "{").replace(r"\}", "}")
    s = s.replace(r"\|", "||")
    s = s.replace(r"\min", "min")
    s = s.replace(r"\ldots", "...")
    s = s.replace(r"\nu", "ν").replace(r"\pi", "π").replace(r"\eta", "η")
    s = s.replace(r"\phi", "φ")
    s = re.sub(r"\\mathrm\{([^{}]+)\}", r"\1", s)
    s = re.sub(r"\\text\{([^{}]+)\}", r"\1", s)
    s = s.replace(r"\sin", "sin").replace(r"\cos", "cos").replace(r"\arccos", "arccos")
    s = re.sub(r"\\sqrt\{([^{}]+)\}", r"sqrt(\1)", s)
    s = re.sub(r"\\frac\{([^{}]+)\}\{([^{}]+)\}", r"(\1)/(\2)", s)
    s = re.sub(r"\\binom\{([^{}]+)\}\{([^{}]+)\}", r"C(\1,\2)", s)
    s = s.replace("{", "").replace("}", "")
    s = s.replace("_\\", "_").replace("^\\", "^")
    s = s.replace(r"\exp", "exp")
    s = re.sub(r"\s+", " ", s)
    # A few common readability cleanups for this manuscript.
    s = s.replace("k_mathrm eff", "k_eff").replace("k_eff", "k_eff")
    s = s.replace("w_mathrm in", "w_in").replace("E_mathrm in", "E_in")
    s = s.replace("s_min", "s_min")
    s = s.replace("LG_p^l'", "LG_p^l'")
    return s


def add_caption(doc: Document, zh: str, en: str) -> None:
    for text, italic in [(zh, False), (en, False)]:
        p = doc.add_paragraph()
        style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=2, line=1.15)
        add_inline_math_runs(p, text, size=9)


def add_figure(doc: Document, num: int, captions: dict[int, tuple[str, str]]) -> None:
    filename, width = FIGURES[num]
    fig_path = FIG_DIR / filename
    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=2, line=1.0)
    run = p.add_run()
    run.add_picture(str(fig_path), width=Cm(width))
    zh, en = captions[num]
    add_caption(doc, zh, en)


def extract_sections(md: str):
    lines = md.splitlines()
    title = ""
    abstract = ""
    keywords = ""
    body_lines: list[str] = []
    fig_lines: list[str] = []
    ref_lines: list[str] = []
    mode = "pre"
    in_abs = False
    for line in lines:
        if line.startswith("# ") and not title:
            title = line[2:].strip()
            continue
        if line.startswith("## 摘要"):
            in_abs = True
            continue
        if in_abs:
            if line.startswith("**关键词**"):
                keywords = line.split("：", 1)[1].strip()
                in_abs = False
                continue
            if line.strip() and not line.startswith("---"):
                abstract += line.strip()
            continue
        if line.startswith("## 图题"):
            mode = "fig"
            continue
        if line.startswith("## 参考文献"):
            mode = "ref"
            continue
        if mode == "fig":
            fig_lines.append(line)
        elif mode == "ref":
            ref_lines.append(line)
        elif mode == "pre":
            if line.startswith("## ") or line.startswith("### ") or line.strip():
                body_lines.append(line)
    return title, abstract, keywords, body_lines, fig_lines, ref_lines


def parse_captions(fig_lines: list[str]) -> dict[int, tuple[str, str]]:
    zh: dict[int, str] = {}
    en: dict[int, str] = {}
    for line in fig_lines:
        line = line.strip()
        m = re.match(r"\*\*图\s*(\d+)\s*(.+?)。\*\*\s*(.*)", line)
        if m:
            n = int(m.group(1))
            zh[n] = f"图 {n} {m.group(2)}。{m.group(3)}".strip()
            continue
        m = re.match(r"\*\*Fig\.\s*(\d+)\s*(.+?)\.\*\*\s*(.*)", line)
        if m:
            n = int(m.group(1))
            en[n] = f"Fig. {n} {m.group(2)}. {m.group(3)}".strip()
    return {n: (zh[n], en[n]) for n in sorted(zh)}


def parse_refs(ref_lines: list[str]) -> list[str]:
    refs = []
    for line in ref_lines:
        line = line.strip()
        if re.match(r"^\[\d+\]", line):
            refs.append(line)
    return refs


def normalize_heading(line: str) -> tuple[int, str] | None:
    if line.startswith("## "):
        text = line[3:].strip()
        if text in {"摘要", "图题", "参考文献"}:
            return None
        return 1, text
    if line.startswith("### "):
        return 2, line[4:].strip()
    return None


def add_front_matter(doc: Document, title: str, abstract: str, keywords: str) -> None:
    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=8, line=1.2)
    add_run(p, title, size=16, bold=True, east="黑体")

    for text in [
        "作者姓名（待补）",
        "（作者单位、城市、邮编待补）",
    ]:
        p = doc.add_paragraph()
        style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, line=1.15)
        add_run(p, text, size=10.5)

    p = doc.add_paragraph()
    style_paragraph(p, first_line=False, after=3, line=1.25)
    add_run(p, "摘  要：", size=10.5, bold=True, east="黑体")
    add_inline_math_runs(p, abstract, size=10.5)

    p = doc.add_paragraph()
    style_paragraph(p, first_line=False, after=3, line=1.25)
    add_run(p, "关键词：", size=10.5, bold=True, east="黑体")
    add_inline_math_runs(p, keywords, size=10.5)

    p = doc.add_paragraph()
    style_paragraph(p, first_line=False, after=8, line=1.25)
    add_run(p, "中图分类号：待补            文献标识码：A", size=10.5)

    p = doc.add_paragraph()
    style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6, line=1.2)
    add_run(p, TITLE_EN, size=14, bold=True, latin="Times New Roman")

    for text in [
        "Author names (to be completed)",
        "(Affiliations to be completed)",
    ]:
        p = doc.add_paragraph()
        style_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, line=1.15)
        add_run(p, text, size=10.5, latin="Times New Roman")

    p = doc.add_paragraph()
    style_paragraph(p, first_line=False, after=3, line=1.25)
    add_run(p, "Abstract: ", size=10.5, bold=True, latin="Times New Roman")
    add_run(p, ABSTRACT_EN, size=10.5, latin="Times New Roman")

    p = doc.add_paragraph()
    style_paragraph(p, first_line=False, after=8, line=1.25)
    add_run(p, "Key words: ", size=10.5, bold=True, latin="Times New Roman")
    add_run(p, KEYWORDS_EN, size=10.5, latin="Times New Roman")


def add_body(doc: Document, body_lines: list[str], captions: dict[int, tuple[str, str]]) -> None:
    inserted_figs: set[int] = set()
    in_eq = False
    eq_lines: list[str] = []
    para_buf: list[str] = []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            text = "".join(x.strip() for x in para_buf).strip()
            para_buf = []
            if text:
                add_plain_paragraph(doc, text, first_line=True)
                for n in range(1, 6):
                    if f"图 {n}" in text and n not in inserted_figs:
                        add_figure(doc, n, captions)
                        inserted_figs.add(n)

    for raw in body_lines:
        line = raw.rstrip()
        if not line.strip() or line.strip() == "---":
            flush_para()
            continue
        if line.strip() == "$$":
            if not in_eq:
                flush_para()
                in_eq = True
                eq_lines = []
            else:
                add_equation(doc, eq_lines)
                in_eq = False
                eq_lines = []
            continue
        if in_eq:
            eq_lines.append(line)
            continue
        h = normalize_heading(line)
        if h:
            flush_para()
            level, text = h
            add_heading(doc, text, level)
            continue
        para_buf.append(line)
    flush_para()


def add_references(doc: Document, refs: list[str]) -> None:
    add_heading(doc, "参考文献", 1)
    for ref in refs:
        p = doc.add_paragraph()
        style_paragraph(p, first_line=False, after=2, line=1.15)
        p.paragraph_format.left_indent = Pt(18)
        p.paragraph_format.first_line_indent = Pt(-18)
        add_run(p, ref, size=9, latin="Times New Roman")


def main() -> None:
    md = MD_PATH.read_text(encoding="utf-8")
    title, abstract, keywords, body_lines, fig_lines, ref_lines = extract_sections(md)
    captions = parse_captions(fig_lines)
    refs = parse_refs(ref_lines)

    doc = Document(DOCX_PATH)
    clear_document(doc)
    set_page(doc)
    add_front_matter(doc, title, abstract, keywords)
    add_body(doc, body_lines, captions)
    add_references(doc, refs)
    doc.save(DOCX_PATH)
    print(f"migrated: {DOCX_PATH}")


if __name__ == "__main__":
    main()
